from __future__ import annotations

import ast
import hashlib
import json
import math
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import brentq
from scipy.stats import kendalltau, linregress, spearmanr

K_B_EV = 8.617333262145e-5
H_EVS = 4.135667696e-15
K_B = 1.380649e-23
H = 6.62607015e-34
AMU = 1.66053906660e-27
BAR = 1e5
EV = 1.602176634e-19
R = 8.314462618
T_COMP = 298.15
MW_NH3_KG_MOL = 0.01703052

ACTIVITY_ORDER_CANON = ['Ru','Os','Fe','Rh','Ir','Co','Re','Mo','Ni','W','Pd','Pt','Cu','Ag','Au']
EN0_CANON = {
    'Ru':-1.1333, 'Os':-1.1095, 'Fe':-1.392120623747696, 'Rh':-0.458, 'Ir':-0.4407,
    'Co':-0.2701, 'Re':-1.8252, 'Mo':-1.8417, 'Ni':-0.0928, 'W':-2.161,
    'Pd':0.5637, 'Pt':0.5895, 'Cu':1.2491, 'Ag':2.911, 'Au':3.0029,
}
MW = {
    'Ru':101.07, 'Os':190.23, 'Fe':55.845, 'Rh':102.9055, 'Ir':192.217,
    'Co':58.933194, 'Re':186.207, 'Mo':95.95, 'Ni':58.6934, 'W':183.84,
    'Pd':106.42, 'Pt':195.084, 'Cu':63.546, 'Ag':107.8682, 'Au':196.96657,
}
PRICE = {
    'Ru':53852.5, 'Os':142650.0, 'Fe':8.0, 'Rh':265243.65919095, 'Ir':252383.36056351,
    'Co':56.28401553583051, 'Re':5757.64, 'Mo':91.22728409215301, 'Ni':16.9094555095895,
    'W':149.97, 'Pd':40670.694409279, 'Pt':51508.71107755406, 'Cu':13.889122517655,
    'Ag':1823.590345370992, 'Au':128732.2322756058,
}
SPECIES = ['H','N','N2','N-N','NH','N-H','NH-H','NH2','NH2-H','NH3']
COMP = {
    'H':(0,1),'N':(1,0),'N2':(2,0),'N-N':(2,0),'NH':(1,1),'N-H':(1,1),
    'NH2':(1,2),'NH-H':(1,2),'NH3':(1,3),'NH2-H':(1,3),
}

_RESPONSE_EGRID = None

def _init_response_worker(egrid):
    global _RESPONSE_EGRID
    _RESPONSE_EGRID = np.asarray(egrid, dtype=float)

def _response_worker(item):
    i, cond = item
    row = np.array([cond.logtof(float(e)) for e in _RESPONSE_EGRID], dtype=float)
    return i, row



def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def frange(start: float, stop: float, step: float) -> np.ndarray:
    return np.round(np.arange(start, stop + step * 0.1, step), 10)


def parse_activity_workbook(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=True, read_only=True)
    required = {'S1_All_444', 'MKM_15', 'Direct_DFT_Check'}
    missing = sorted(required - set(wb.sheetnames))
    if missing:
        raise RuntimeError(f"Missing required activity workbook sheets: {missing}")
    ws = wb['S1_All_444']
    rows = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        surface, phase, site, species, energy, freqs, ref, coverage, std = r[:9]
        if species is None or energy is None:
            continue
        try:
            fs = [float(x) for x in ast.literal_eval(freqs)] if isinstance(freqs, str) else []
        except Exception:
            fs = []
        rows.append({'surface':surface,'phase':phase,'site':site,'species':species,'E':float(energy),'freqs':fs})
    D = {(r['phase'], r['surface'], r['species']): r for r in rows}
    phase = 'step_site'
    metals_source = sorted({r['surface'] for r in rows if r['surface'] is not None})
    avail = [m for m in metals_source if (phase,m,'N') in D]
    if len(avail) != 14:
        raise RuntimeError(f"Expected 14 raw step-site metals in S1 workbook; found {len(avail)}")
    fits = {}
    for sp in SPECIES:
        xs = [D[(phase,m,'N')]['E'] for m in avail]
        ys = [D[(phase,m,sp)]['E'] for m in avail]
        lr = linregress(xs, ys)
        fits[sp] = (float(lr.slope), float(lr.intercept))
    median_freqs = {}
    for sp in SPECIES:
        lists = [D[(phase,m,sp)]['freqs'] for m in avail if D[(phase,m,sp)]['freqs']]
        nmode = min(map(len, lists)) if lists else 0
        median_freqs[sp] = [float(np.median([fs[i] for fs in lists])) for i in range(nmode)]
    gasrec = {}
    for r in rows:
        if r['surface'] is None and r['species'] in ('H2','N2','NH3'):
            gasrec.setdefault(r['species'], r)
    if set(gasrec) != {'H2','N2','NH3'}:
        raise RuntimeError(f"Gas records incomplete: {sorted(gasrec)}")

    # Extract the workbook's own headline activity order as an independent gate.
    activity_sheet = wb['MKM_15']
    workbook_order = []
    for r in activity_sheet.iter_rows(min_row=5, values_only=True):
        if r[1] is not None:
            workbook_order.append(str(r[1]))
    direct_sheet = wb['Direct_DFT_Check']
    direct_ru_tof = None
    for r in direct_sheet.iter_rows(min_row=5, values_only=True):
        if r[1] == 'Ru':
            direct_ru_tof = float(r[4]); break
    scaling_ru_tof = None
    for r in activity_sheet.iter_rows(min_row=5, values_only=True):
        if r[1] == 'Ru':
            scaling_ru_tof = float(r[4]); break

    return {
        'rows': rows, 'D': D, 'fits': fits, 'median_freqs': median_freqs, 'gasrec': gasrec,
        'workbook_activity_order': workbook_order,
        'direct_ru_tof': direct_ru_tof, 'scaling_ru_tof': scaling_ru_tof,
    }


def vib_F(freqs_meV, T):
    kT = K_B_EV * T
    out = 0.0
    for f in freqs_meV:
        eps = f / 1000.0
        if eps <= 1e-12:
            continue
        x = eps / kT
        out += 0.5 * eps if x > 700 else 0.5 * eps + kT * math.log1p(-math.exp(-x))
    return out


def gas_corr_linear(mass_amu, I, sigma, freqs_meV, T, p0=BAR):
    m = mass_amu * AMU
    qtrans = (2 * math.pi * m * K_B * T / H**2)**1.5 * (K_B*T/p0)
    Strans = K_B * (math.log(qtrans) + 2.5)
    Htrans = 2.5 * K_B * T
    qrot = 8 * math.pi**2 * I * K_B * T / (sigma * H**2)
    Srot = K_B * (math.log(qrot) + 1.0)
    Hrot = K_B * T
    Hv = 0.0; Sv = 0.0
    for f in freqs_meV:
        epsJ = f/1000.0*EV; x = epsJ/(K_B*T)
        Hv += 0.5*epsJ + epsJ/(math.exp(x)-1.0)
        Sv += K_B*(x/(math.exp(x)-1.0) - math.log1p(-math.exp(-x)))
    return ((Htrans+Hrot+Hv) - T*(Strans+Srot+Sv))/EV


def gas_corr_nonlinear(mass_amu, moments, sigma, freqs_meV, T, p0=BAR):
    m = mass_amu * AMU
    qtrans = (2 * math.pi * m * K_B * T / H**2)**1.5 * (K_B*T/p0)
    Strans = K_B*(math.log(qtrans)+2.5); Htrans = 2.5*K_B*T
    IA,IB,IC = moments
    qrot = math.sqrt(math.pi)/sigma*(8*math.pi**2*K_B*T/H**2)**1.5*math.sqrt(IA*IB*IC)
    Srot = K_B*(math.log(qrot)+1.5); Hrot = 1.5*K_B*T
    Hv = 0.0; Sv = 0.0
    for f in freqs_meV:
        epsJ = f/1000.0*EV; x = epsJ/(K_B*T)
        Hv += 0.5*epsJ + epsJ/(math.exp(x)-1.0)
        Sv += K_B*(x/(math.exp(x)-1.0) - math.log1p(-math.exp(-x)))
    return ((Htrans+Hrot+Hv) - T*(Strans+Srot+Sv))/EV


def nh3_moments():
    rH2=0.7414e-10; rN2=1.0977e-10
    I_H2=(1.00784*AMU)*(rH2**2)/2.0
    I_N2=(14.0067*AMU)*(rN2**2)/2.0
    NH=1.012e-10; gamma=math.radians(106.7)
    cos2=(math.cos(gamma)+0.5)/1.5; alpha=math.acos(math.sqrt(max(cos2,0)))
    coords=[np.array([0.,0.,0.])]; masses=[14.0067*AMU]
    for phi in [0,2*math.pi/3,4*math.pi/3]:
        coords.append(np.array([NH*math.sin(alpha)*math.cos(phi),NH*math.sin(alpha)*math.sin(phi),NH*math.cos(alpha)]))
        masses.append(1.00784*AMU)
    M=sum(masses); com=sum(m*c for m,c in zip(masses,coords))/M
    I=np.zeros((3,3))
    for m,c in zip(masses,coords):
        rr=c-com; I += m*((rr@rr)*np.eye(3)-np.outer(rr,rr))
    return I_H2, I_N2, np.linalg.eigvalsh(I)


I_H2, I_N2, NH3_MOMENTS = nh3_moments()


@dataclass
class KineticContext:
    fits: dict
    median_freqs: dict
    gasrec: dict

    def gas_G(self, T):
        out = {}
        for sp in ('H2','N2','NH3'):
            r = self.gasrec[sp]
            if sp == 'H2': corr = gas_corr_linear(2.01588, I_H2, 2, r['freqs'], T)
            elif sp == 'N2': corr = gas_corr_linear(28.0134, I_N2, 2, r['freqs'], T)
            else: corr = gas_corr_nonlinear(17.03052, NH3_MOMENTS, 3, r['freqs'], T)
            out[sp] = r['E'] + corr
        return out

    def condition(self, T, pN2, pH2, pNH3):
        return Condition(self, T, pN2, pH2, pNH3)


class Condition:
    def __init__(self, ctx: KineticContext, T, pN2, pH2, pNH3):
        self.ctx = ctx; self.T=float(T); self.kT=K_B_EV*self.T; self.pref=self.kT/H_EVS
        self.pN2=float(pN2); self.pH2=float(pH2); self.pNH3=float(pNH3)
        gg=ctx.gas_G(self.T)
        self.G_NH3_g_form=gg['NH3']-0.5*gg['N2']-1.5*gg['H2']
        self.g_linear={}
        for sp in SPECIES:
            slope,intercept=ctx.fits[sp]; x,y=COMP[sp]
            a=1.0 if sp=='N' else slope
            b=(0.0 if sp=='N' else intercept)+vib_F(ctx.median_freqs[sp],self.T)-x/2*gg['N2']-y/2*gg['H2']
            self.g_linear[sp]=(a,b)

    @staticmethod
    def _exp(x): return math.exp(max(-745.0,min(700.0,x)))

    def logtof(self, EN):
        kT=self.kT; pref=self.pref; g={sp:a*EN+b for sp,(a,b) in self.g_linear.items()}
        K_N2=self._exp(-g['N2']/kT); K_H=self._exp(-g['H']/kT); K_NH3=self._exp(-(g['NH3']-self.G_NH3_g_form)/kT)
        qN2=K_N2*self.pN2; qH=K_H*math.sqrt(self.pH2); qNH3=K_NH3*self.pNH3
        def kfkr(ts,react_sum,prod_sum):
            gts=max(g[ts],react_sum,prod_sum)
            return pref*self._exp(-(gts-react_sum)/kT), pref*self._exp(-(gts-prod_sum)/kT)
        k2f,k2r=kfkr('N-N',g['N2'],2*g['N']); k4f,k4r=kfkr('N-H',g['N']+g['H'],g['NH'])
        k5f,k5r=kfkr('NH-H',g['NH']+g['H'],g['NH2']); k6f,k6r=kfkr('NH2-H',g['NH2']+g['H'],g['NH3'])
        tiny=1e-300
        def ratios_from_j(j):
            z=(j+k6r*qNH3)/max(k6f*qH,tiny); y=(j+k5r*z)/max(k5f*qH,tiny); x=(j+k4r*y)/max(k4f*qH,tiny); return x,y,z
        A=max(k2f*qN2,tiny); B=k2r; jhi=2.0*A
        def f_scaled(u):
            x,_,_=ratios_from_j(jhi*u); return u-1.0+(B/A)*x*x
        if f_scaled(0.0)<=0.0 and f_scaled(1.0)>=0.0:
            u=brentq(f_scaled,0.0,1.0,xtol=1e-14,rtol=1e-13,maxiter=200)
        else:
            ulo=-1.0
            for _ in range(80):
                if f_scaled(ulo)*f_scaled(0.0)<=0.0: break
                ulo*=10.0
            u=brentq(f_scaled,ulo,0.0,xtol=1e-14,rtol=1e-13,maxiter=200)
        j=jhi*u; x,y,z=ratios_from_j(j); x=max(x,0.0); y=max(y,0.0); z=max(z,0.0)
        theta_star=1.0/(1+qN2+qH+qNH3+x+y+z); tof=j*theta_star**2
        return math.log10(tof) if tof>0 else -300.0


class NH3Harness:
    def __init__(self, cfg: dict, root: Path):
        self.cfg=cfg; self.root=root
        inp=root/cfg['inputs']['activity_workbook']
        self.input_path=inp
        src=parse_activity_workbook(inp)
        self.src=src; self.ctx=KineticContext(src['fits'],src['median_freqs'],src['gasrec'])
        self.activity_order=ACTIVITY_ORDER_CANON.copy(); self.EN0=EN0_CANON.copy()
        self._init_controls()
        self.process_states=self._build_process_states()
        self._state_arrays()
        self.base_condition=self.ctx.condition(cfg['atomic_reference']['temperature_K'],cfg['atomic_reference']['pN2_bar'],cfg['atomic_reference']['pH2_bar'],cfg['atomic_reference']['pNH3_bar'])
        ng=cfg['numerics']['descriptor_grid_eV']; self.EGRID=np.round(np.arange(float(ng['start']), float(ng['stop']) + 1e-9, float(ng['step'])), 3)
        self.response_base=np.array([self.base_condition.logtof(float(e)) for e in self.EGRID])

    def _init_controls(self):
        c=self.cfg; proc=c['process']; eng=c['engineering']; plant=c['plant']; econ=c['economics']; unc=c['uncertainty']
        self.T_GRID_C=[float(x) for x in proc['temperature_C']]
        self.P_GRID_BAR=frange(proc['pressure_bar']['start'],proc['pressure_bar']['stop'],proc['pressure_bar']['step'])
        self.TSEP_GRID_C=frange(proc['separator_temperature_C']['start'],proc['separator_temperature_C']['stop'],proc['separator_temperature_C']['step'])
        self.APPROACH_EQ=float(proc['approach_to_equilibrium']); self.DP_RECYCLE_BAR=float(proc['recycle_pressure_rise_bar'])
        self.RECYCLE_NH3_CAP=float(proc['max_separator_vapor_nh3_fraction']); self.REFRIG_ETA=float(proc['refrigeration_efficiency_fraction_carnot'])
        self.REFRIG_SINK_C=float(proc['refrigeration_sink_C']); self.DHVAP=float(proc['ammonia_heat_of_vaporization_kJ_mol'])
        self.V_CAP=float(eng['max_catalyst_bed_m3']); self.BED_DENSITY=float(eng['bed_density_kg_m3']); self.ACTIVE_FRACTION=float(eng['active_fraction'])
        self.PLANT_TPD=float(plant['production_tpd']); self.CAPACITY_FACTOR=float(plant['capacity_factor']); self.FRESH_FEED_PRESSURE_BAR=float(plant['fresh_feed_pressure_bar'])
        self.DISCOUNT_RATE=float(econ['discount_rate']); self.PLANT_LIFE_Y=int(econ['plant_life_y']); self.ELECTRICITY=float(econ['electricity_USD_MWh'])
        self.COMPRESSOR_ETA=float(econ['compressor_efficiency']); self.CATALYST_LIFE_Y=float(econ['catalyst_life_y']); self.METAL_RECOVERY=float(econ['metal_recovery_fraction'])
        self.REACTOR_FIXED=float(econ['reactor_fixed_USD']); self.REACTOR_VAR=float(econ['reactor_variable_USD']); self.REACTOR_REF=float(econ['reactor_reference_m3']); self.REACTOR_EXP=float(econ['reactor_exponent']); self.F_CAL=float(econ['calibration_factor'])
        self.MC_N=int(unc['draws']); self.MC_SEED=int(unc['seed']); self.FE_SIGMA=float(unc['Fe_sigma_eV']); self.OTHER_HALF_WIDTH=float(unc['other_uniform_half_width_eV'])
        # Optional pressure-dependent CAPEX (NH3-FINAL-1.1). Absent or enabled=false -> bit-identical to NH3-FINAL-1.0.
        pc=econ.get('pressure_capex') or {}
        self.PCAPEX_ENABLED=bool(pc.get('enabled',False))
        if self.PCAPEX_ENABLED:
            self.PC_INDEX_RATIO=float(pc['cost_index_current'])/float(pc['cost_index_base'])
            v=pc['vessel']; self.PC_VK=[float(x) for x in v['K']]; self.PC_V_AMIN=float(v['A_min_m3']); self.PC_V_AMAX=float(v['A_max_m3']); self.PC_V_NEXP=float(v['extrapolation_exponent'])
            self.PC_V_B2=float(v['B2']); self.PC_V_FM=float(v['Fm']); self.PC_V_LD=float(v['L_over_D']); self.PC_V_SE=float(v['allowable_stress_x_weld_eff_bar']); self.PC_V_CA=float(v['corrosion_allowance_m']); self.PC_V_TMIN=float(v['min_thickness_m'])
            c=pc['compressor']; self.PC_CK=[float(x) for x in c['K']]; self.PC_C_AMIN=float(c['A_min_kW']); self.PC_C_AMAX=float(c['A_max_kW']); self.PC_C_NEXP=float(c['extrapolation_exponent']); self.PC_C_FBM=float(c['F_BM'])
        self.annual_output_t=self.PLANT_TPD*365.0*self.CAPACITY_FACTOR
        self.nh3_mol_s=self.PLANT_TPD*1000.0/MW_NH3_KG_MOL/86400.0
        self.crf=self.DISCOUNT_RATE*(1+self.DISCOUNT_RATE)**self.PLANT_LIFE_Y/((1+self.DISCOUNT_RATE)**self.PLANT_LIFE_Y-1)
        self.n_per_t=1000.0/MW_NH3_KG_MOL

    def psat_nh3_bar(self,T_C):
        T=T_C+273.15
        if T<239.6: A,B,C=3.18757,506.713,-80.78
        else: A,B,C=4.86886,1113.928,-10.409
        return 10.0**(A-B/(T+C))

    def refriger_cost(self,Tsep_C):
        if Tsep_C>=self.REFRIG_SINK_C: return 0.0
        Tc=Tsep_C+273.15; Th=self.REFRIG_SINK_C+273.15
        Wrev=self.DHVAP*1000.0*(Th-Tc)/Tc; Wact=Wrev/self.REFRIG_ETA
        return Wact*self.n_per_t/3.6e9*self.ELECTRICITY

    # ---- pressure-dependent CAPEX helpers (NH3-FINAL-1.1; all return 0 when disabled) ----
    @staticmethod
    def turton_cp0(K,A,Amin,Amax,nexp):
        """Turton purchased cost at the correlation basis index: log10 Cp0 = K1 + K2 log10 A + K3 (log10 A)^2 inside [Amin, Amax];
        outside the range the cost is extrapolated from the nearest bound with a power law of exponent nexp."""
        A=np.asarray(A,dtype=float); Ac=np.clip(A,Amin,Amax); la=np.log10(Ac)
        return np.power(10.0,K[0]+K[1]*la+K[2]*la*la)*np.power(A/Ac,nexp)

    def compressor_capex_USD_t(self,elec_cost_USD_t):
        """Annualized bare-module cost (USD per t NH3) of one compressor service sized from its electricity cost at design capacity.
        Shaft power kW = (USD/t / USD/MWh * 1000 kWh/MWh) * tpd / 24; services above A_max are split into equal units."""
        if not self.PCAPEX_ENABLED or elec_cost_USD_t<=0.0: return 0.0
        W_kW=elec_cost_USD_t/self.ELECTRICITY*1000.0*self.PLANT_TPD/24.0
        n=max(1,int(math.ceil(W_kW/self.PC_C_AMAX))); Wu=W_kW/n
        cp0=n*float(self.turton_cp0(self.PC_CK,Wu,self.PC_C_AMIN,self.PC_C_AMAX,self.PC_C_NEXP))
        return cp0*self.PC_C_FBM*self.PC_INDEX_RATIO*self.crf/self.annual_output_t

    def vessel_pressure_premium(self,V):
        """Pressure-attributable converter-shell cost (USD/t) per process state: Turton vertical-vessel Cp0(V) * B2 * Fm * (Fp - 1),
        Fp = max(t/t_min, 1) with ASME wall thickness t = P D / (2 (S E - 0.6 P)) + CA, D from V at fixed L/D.
        Vessels below A_min are costed at A_min. Zero when disabled (1.0 behaviour)."""
        V=np.asarray(V,dtype=float)
        if not self.PCAPEX_ENABLED: return np.zeros_like(V)
        Vd=np.maximum(V,self.PC_V_AMIN); D=np.power(4.0*Vd/(math.pi*self.PC_V_LD),1.0/3.0); P=self.state_P
        denom=2.0*(self.PC_V_SE-0.6*P)
        t=np.where(denom>0.0,P*D/np.where(denom>0.0,denom,1.0)+self.PC_V_CA,np.inf)
        Fp=np.maximum(t/self.PC_V_TMIN,1.0)
        cp0=self.turton_cp0(self.PC_VK,Vd,self.PC_V_AMIN,self.PC_V_AMAX,self.PC_V_NEXP)
        return cp0*self.PC_V_B2*self.PC_V_FM*(Fp-1.0)*self.PC_INDEX_RATIO*self.crf/self.annual_output_t

    def cost_breakdown_at(self,metal,V,metal_cost,reactor,i):
        """Cost pools (USD/t) at state i for reporting; reactor = base volume proxy + vessel pressure premium."""
        Vi=float(V[i]); base=float((self.REACTOR_FIXED+self.REACTOR_VAR*(Vi/self.REACTOR_REF)**self.REACTOR_EXP)*self.crf/self.annual_output_t)
        return dict(metal_cost=float(metal_cost[i]),reactor_base=base,vessel_pressure_premium=float(reactor[i])-base,fresh_comp=float(self.state_fresh[i]),recycle_comp=float(self.state_reccomp[i]),refrigeration=float(self.state_refrig[i]),compressor_capex=float(self.state_compcapex[i]))

    @staticmethod
    def equilibrium_extent(Kp,P,nN,nH,nA):
        def Q(xi):
            N=nN-xi; Hh=nH-3.0*xi; A=nA+2.0*xi; tot=N+Hh+A
            return (A/tot*P)**2/((N/tot*P)*(Hh/tot*P)**3)
        hi=min(nN,nH/3.0)*(1.0-1e-10)
        if Q(0.0)>=Kp: return 0.0
        return brentq(lambda z:Q(z)-Kp,0.0,hi,xtol=1e-12,rtol=1e-10,maxiter=100)

    def self_consistent_process_state(self,C,P,Tsep):
        T=C+273.15; psat=self.psat_nh3_bar(Tsep)
        if psat>=P: return None
        ysep=psat/P
        if ysep>self.RECYCLE_NH3_CAP: return None
        gg=self.ctx.gas_G(T); Gform=gg['NH3']-0.5*gg['N2']-1.5*gg['H2']; Kp=math.exp(-2.0*Gform/(K_B_EV*T)); ratio=ysep/(1.0-ysep)
        def fixed_point(X):
            nN=0.5/X; nH=1.5/X; nA=ratio*2.0*(1.0-X)/X
            xi_eq=self.equilibrium_extent(Kp,P,nN,nH,nA); return X-self.APPROACH_EQ*(xi_eq/nN)
        lo,hi=1e-6,0.999; flo,fhi=fixed_point(lo),fixed_point(hi)
        if flo>=0.0 or fhi<=0.0: return None
        X=brentq(fixed_point,lo,hi,xtol=1e-10,rtol=1e-9,maxiter=100); Xeq=X/self.APPROACH_EQ
        reactive_recycle=2.0*(1.0-X)/X; nh3_recycle=ratio*reactive_recycle; recycle_total=reactive_recycle+nh3_recycle
        reactive_in=2.0/X; total_in=reactive_in+nh3_recycle
        yN=(0.5/X)/total_in; yH=(1.5/X)/total_in; yA=nh3_recycle/total_in
        pN2,pH2,pNH3=P*yN,P*yH,max(P*yA,1e-12); cond=self.ctx.condition(T,pN2,pH2,pNH3)
        fresh=0.0 if P<=self.FRESH_FEED_PRESSURE_BAR else (2.0*self.n_per_t*R*T_COMP*math.log(P/self.FRESH_FEED_PRESSURE_BAR)/self.COMPRESSOR_ETA/3.6e9*self.ELECTRICITY)
        if P<=self.DP_RECYCLE_BAR: return None
        recycle_comp=self.n_per_t*recycle_total*R*T_COMP*math.log(P/(P-self.DP_RECYCLE_BAR))/self.COMPRESSOR_ETA/3.6e9*self.ELECTRICITY
        refrig=self.refriger_cost(Tsep); inlet_Nm3_h=self.nh3_mol_s*total_in*0.022414*3600.0
        comp_capex=self.compressor_capex_USD_t(fresh)+self.compressor_capex_USD_t(recycle_comp)
        return dict(T_C=float(C),P_bar=float(P),Tsep_C=float(Tsep),X=float(X),Xeq=float(Xeq),ysep=float(ysep),recycle_total=float(recycle_total),inlet_Nm3_h=float(inlet_Nm3_h),fresh_comp=float(fresh),recycle_comp=float(recycle_comp),refrigeration=float(refrig),compressor_capex=float(comp_capex),process_cost=float(fresh+recycle_comp+refrig+comp_capex),condition=cond)

    def _build_process_states(self):
        out=[]
        for C in self.T_GRID_C:
            for P in self.P_GRID_BAR:
                for Tsep in self.TSEP_GRID_C:
                    s=self.self_consistent_process_state(float(C),float(P),float(Tsep))
                    if s is not None: out.append(s)
        return out

    def _state_arrays(self):
        ss=self.process_states
        self.state_T=np.array([s['T_C'] for s in ss]); self.state_P=np.array([s['P_bar'] for s in ss]); self.state_Tsep=np.array([s['Tsep_C'] for s in ss])
        self.state_process_cost=np.array([s['process_cost'] for s in ss]); self.state_fresh=np.array([s['fresh_comp'] for s in ss]); self.state_reccomp=np.array([s['recycle_comp'] for s in ss]); self.state_refrig=np.array([s['refrigeration'] for s in ss]); self.state_compcapex=np.array([s['compressor_capex'] for s in ss])
        self.NSTATE=len(ss)

    def direct_state_logtof(self,EN):
        return np.array([s['condition'].logtof(float(EN)) for s in self.process_states],dtype=float)

    def frozen_state_logtof(self, EN):
        """Reproduce the frozen NH3-FINAL-1.0 descriptor interpolation exactly.

        The published closure evaluates each process state's MKM on the 0.005-eV
        descriptor grid and linearly interpolates to the candidate E_N.  Computing
        only the two bracketing grid points gives the identical baseline result
        without constructing the full 3636 x 1201 response matrix in smoke mode.
        """
        x=min(max(float(EN),float(self.EGRID[0])),float(self.EGRID[-1]))
        pos=(x-float(self.EGRID[0]))/float(self.EGRID[1]-self.EGRID[0])
        j=int(math.floor(pos))
        if j>=len(self.EGRID)-1:
            e0=e1=float(self.EGRID[-1]); f=0.0
        else:
            e0=float(self.EGRID[j]); e1=float(self.EGRID[j+1]); f=pos-j
        out=np.empty(self.NSTATE,dtype=float)
        if f==0.0 or e0==e1:
            for i,s in enumerate(self.process_states): out[i]=s['condition'].logtof(e0)
        else:
            for i,s in enumerate(self.process_states):
                a=s['condition'].logtof(e0); b=s['condition'].logtof(e1); out[i]=a*(1.0-f)+b*f
        return out

    def cost_arrays(self,metal,logtof,price=None,life_y=None,recovery=None,alpha=1.0):
        if price is None: price=PRICE[metal]
        if life_y is None: life_y=self.CATALYST_LIFE_Y
        if recovery is None: recovery=self.METAL_RECOVERY
        active=self.nh3_mol_s*MW[metal]/self.F_CAL*np.power(10.0,-logtof)/alpha
        V=active/self.ACTIVE_FRACTION/self.BED_DENSITY
        metal_cost=active*price*(1.0-recovery)/life_y/self.annual_output_t
        reactor=(self.REACTOR_FIXED+self.REACTOR_VAR*np.power(V/self.REACTOR_REF,self.REACTOR_EXP))*self.crf/self.annual_output_t+self.vessel_pressure_premium(V)
        total=metal_cost+reactor+self.state_process_cost
        return total,V,metal_cost,reactor

    def deterministic(self):
        activity_score={m:self.base_condition.logtof(self.EN0[m]) for m in self.activity_order}
        activity_order=sorted(self.activity_order,key=lambda m:activity_score[m],reverse=True); activity_rank={m:i+1 for i,m in enumerate(activity_order)}
        det={}
        for m in self.activity_order:
            logvec=self.frozen_state_logtof(self.EN0[m]); total,V,metal_cost,reactor=self.cost_arrays(m,logvec)
            iu=int(np.argmin(total)); eligible=V<=self.V_CAP
            feas=None
            if np.any(eligible):
                im=int(np.argmin(np.where(eligible,total,np.inf)))
                feas=dict(total_cost=float(total[im]),V_m3=float(V[im]),logTOF=float(logvec[im]),T_C=float(self.state_T[im]),P_bar=float(self.state_P[im]),Tsep_C=float(self.state_Tsep[im]),state_index=im,breakdown=self.cost_breakdown_at(m,V,metal_cost,reactor,im))
            det[m]=dict(activity_logTOF=float(activity_score[m]),min_bed_m3=float(np.min(V)),unconstrained=dict(total_cost=float(total[iu]),V_m3=float(V[iu]),T_C=float(self.state_T[iu]),P_bar=float(self.state_P[iu]),Tsep_C=float(self.state_Tsep[iu]),breakdown=self.cost_breakdown_at(m,V,metal_cost,reactor,iu)),feasible=feas)
        raw_order=sorted(self.activity_order,key=lambda m:det[m]['unconstrained']['total_cost']); raw_rank={m:i+1 for i,m in enumerate(raw_order)}
        fmet=[m for m in self.activity_order if det[m]['feasible'] is not None]; ford=sorted(fmet,key=lambda m:det[m]['feasible']['total_cost']); frank={m:i+1 for i,m in enumerate(ford)}
        nf=len(ford); tie=(nf+1+len(self.activity_order))/2.0 if nf<len(self.activity_order) else None; cens={m:(frank[m] if m in frank else tie) for m in self.activity_order}
        raw_rho=float(spearmanr([activity_rank[m] for m in self.activity_order],[raw_rank[m] for m in self.activity_order]).statistic)
        cens_rho=float(spearmanr([activity_rank[m] for m in self.activity_order],[cens[m] for m in self.activity_order]).statistic)
        rolling=[]
        for K in range(3,len(self.activity_order)+1):
            subset=activity_order[:K]; ar={m:i+1 for i,m in enumerate(subset)}; eraw_order=sorted(subset,key=lambda m:det[m]['unconstrained']['total_cost']); eraw={m:i+1 for i,m in enumerate(eraw_order)}
            rho_raw=float(spearmanr([ar[m] for m in subset],[eraw[m] for m in subset]).statistic)
            fset=[m for m in subset if det[m]['feasible'] is not None]; fo=sorted(fset,key=lambda m:det[m]['feasible']['total_cost']); fr={m:i+1 for i,m in enumerate(fo)}; nfk=len(fo); tie_k=(nfk+1+K)/2.0 if nfk<K else None; cr=[fr[m] if m in fr else tie_k for m in subset]
            rho_c=float(spearmanr([ar[m] for m in subset],cr).statistic) if len(set(cr))>1 else math.nan
            rolling.append({'K':K,'rho_raw':rho_raw,'rho_censored':rho_c,'n_feasible':nfk})
        return {'activity_order':activity_order,'raw_economic_order':raw_order,'feasible_economic_order':ford,'raw_global_spearman':raw_rho,'censored_global_spearman':cens_rho,'rolling':rolling,'metals':det}

    def response_cache_path(self):
        sig=sha256_file(self.input_path)[:16]; ng=self.cfg['numerics']['descriptor_grid_eV']
        return self.root/'cache'/f"response_{sig}_{ng['start']}_{ng['stop']}_{ng['step']}_{self.NSTATE}.npz"

    def build_or_load_response(self,force=False):
        cp=self.response_cache_path(); cp.parent.mkdir(parents=True,exist_ok=True)
        if cp.exists() and not force:
            z=np.load(cp); resp=z['response']; eg=z['EGRID']
            if resp.shape==(self.NSTATE,len(self.EGRID)) and np.allclose(eg,self.EGRID): return resp,cp,True
        resp=np.empty((self.NSTATE,len(self.EGRID)),dtype=float)
        workers=max(1,int(self.cfg.get('numerics',{}).get('response_workers',1)))
        if workers == 1:
            for i,s in enumerate(self.process_states):
                cond=s['condition']; resp[i,:]=[cond.logtof(float(e)) for e in self.EGRID]
                if (i+1)%250==0: print(f"  response states: {i+1}/{self.NSTATE}",flush=True)
        else:
            from multiprocessing import Pool
            items=((i,s['condition']) for i,s in enumerate(self.process_states))
            done=0
            with Pool(processes=workers, initializer=_init_response_worker, initargs=(self.EGRID,)) as pool:
                for i,row in pool.imap_unordered(_response_worker, items, chunksize=8):
                    resp[i,:]=row; done+=1
                    if done%250==0: print(f"  response states: {done}/{self.NSTATE}",flush=True)
        np.savez_compressed(cp,response=resp,EGRID=self.EGRID); return resp,cp,False

    def interp_state_vector(self,response,x):
        x=min(max(float(x),float(self.EGRID[0])),float(self.EGRID[-1])); pos=(x-self.EGRID[0])/(self.EGRID[1]-self.EGRID[0]); j=int(math.floor(pos))
        if j>=len(self.EGRID)-1: return response[:,-1].copy()
        f=pos-j; return response[:,j]*(1-f)+response[:,j+1]*f

    def monte_carlo(self,response):
        rng=np.random.default_rng(self.MC_SEED); M=len(self.activity_order); ens=np.zeros((self.MC_N,M),dtype=float)
        for d in range(self.MC_N):
            for j,m in enumerate(self.activity_order):
                ens[d,j]=rng.normal(self.EN0[m],self.FE_SIGMA) if m=='Fe' else rng.uniform(self.EN0[m]-self.OTHER_HALF_WIDTH,self.EN0[m]+self.OTHER_HALF_WIDTH)
        activity_mc=np.empty_like(ens)
        for j,m in enumerate(self.activity_order): activity_mc[:,j]=np.interp(ens[:,j],self.EGRID,self.response_base)
        raw_rhos=[]; cens_rhos=[]; top1=[]; top3a=[]; top3c=[]; nfeas=[]; fe_feas=[]
        for d in range(self.MC_N):
            act_order_idx=np.argsort(-activity_mc[d]); arank=np.empty(M,int); arank[act_order_idx]=np.arange(1,M+1)
            raw_cost=np.empty(M); feas_cost=np.full(M,np.inf)
            for j,m in enumerate(self.activity_order):
                logv=self.interp_state_vector(response,ens[d,j]); total,V,_,_=self.cost_arrays(m,logv); raw_cost[j]=float(np.min(total)); mask=V<=self.V_CAP
                if np.any(mask): feas_cost[j]=float(np.min(np.where(mask,total,np.inf)))
            raw_order_idx=np.argsort(raw_cost); rrank=np.empty(M,int); rrank[raw_order_idx]=np.arange(1,M+1)
            raw_rhos.append(float(spearmanr(arank,rrank).statistic))
            fidx=np.where(np.isfinite(feas_cost))[0]; nf=len(fidx); nfeas.append(nf); fe_feas.append(bool(np.isfinite(feas_cost[self.activity_order.index('Fe')])))
            if nf:
                fsort=fidx[np.argsort(feas_cost[fidx])]; frank={int(idx):i+1 for i,idx in enumerate(fsort)}; tie=(nf+1+M)/2.0 if nf<M else None; cr=np.array([frank[i] if i in frank else tie for i in range(M)],float)
                cens_rhos.append(float(spearmanr(arank,cr).statistic) if len(set(cr))>1 else math.nan)
                top1.append(float(act_order_idx[0]==fsort[0]))
                atop3=set(act_order_idx[:3]); ftop3=set(fsort[:3]); inter=len(atop3 & ftop3); top3a.append(inter/3.0); top3c.append(inter/min(3,nf))
            else:
                cens_rhos.append(math.nan); top1.append(0.0); top3a.append(0.0); top3c.append(math.nan)
        return {'raw_rho_mean':float(np.mean(raw_rhos)),'censored_rho_mean':float(np.nanmean(cens_rhos)),'top1_survival':float(np.mean(top1)),'top3_actionable':float(np.mean(top3a)),'top3_conditional':float(np.nanmean(top3c)),'mean_feasible_count':float(np.mean(nfeas)),'Fe_feasibility_probability':float(np.mean(fe_feas))}

    def eval_scenario(self,metal,base_logvec,alpha=1.0,price=None,life_y=None,recovery=None):
        total,V,metal_cost,reactor=self.cost_arrays(metal,base_logvec,price=price,life_y=life_y,recovery=recovery,alpha=alpha); mask=V<=self.V_CAP
        if not np.any(mask): raise RuntimeError('No feasible state')
        i=int(np.argmin(np.where(mask,total,np.inf)))
        return {'total_cost':float(total[i]),'V_m3':float(V[i]),'T_C':float(self.state_T[i]),'P_bar':float(self.state_P[i]),'Tsep_C':float(self.state_Tsep[i]),'logTOF':float(base_logvec[i]+math.log10(alpha)),'state_index':i,'breakdown':self.cost_breakdown_at(metal,V,metal_cost,reactor,i)}

    def backward_and_reachability(self,response):
        fe_log=self.interp_state_vector(response,self.EN0['Fe']); ru_log=self.interp_state_vector(response,self.EN0['Ru']); fe=self.eval_scenario('Fe',fe_log); FE=fe['total_cost']
        def gap(loga):
            try: return self.eval_scenario('Ru',ru_log,alpha=10.0**loga)['total_cost']-FE
            except RuntimeError: return 1e100
        alpha_star=10.0**brentq(gap,-6.0,12.0,xtol=1e-11)
        base_ru_ref=float(self.base_condition.logtof(self.EN0['Ru'])); max_ref=float(np.max(self.response_base)); ref_gain=10.0**(max_ref-base_ru_ref)
        ru_state=self.interp_state_vector(response,self.EN0['Ru']); state_gain=float(np.max(np.power(10.0,np.max(response,axis=1)-ru_state)))
        # Strict scaling cost search on the same process library, as in closure: E_N -2.2 to 0.2 eV.
        maskE=(self.EGRID>=-2.2)&(self.EGRID<=0.2); best=math.inf; best_EN=None
        for j in np.where(maskE)[0]:
            total,V,_,_=self.cost_arrays('Ru',response[:,j]); ok=V<=self.V_CAP
            if np.any(ok):
                val=float(np.min(np.where(ok,total,np.inf)))
                if val<best: best=val; best_EN=float(self.EGRID[j])
        direct_gain=None
        if self.src['direct_ru_tof'] and self.src['scaling_ru_tof']:
            direct_gain=float(self.src['direct_ru_tof']/self.src['scaling_ru_tof'])
        return {'Ru_activity_break_even_multiplier':float(alpha_star),'Ru_scaling_max_gain_673K':float(ref_gain),'Ru_scaling_max_gain_all_states':state_gain,'Ru_best_scaling_cost_USD_t':float(best),'Ru_best_scaling_EN_eV':best_EN,'direct_DFT_diagnostic_gain':direct_gain}
