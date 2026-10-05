"""Validation of the generalized methanol recycle-economics model (`data/meoh/meoh_general_model.py`).

1. Canonical reproduction: with x_CO = 0, P = 100 bar, H2/CO2 = 4 and the STY-per-g-Re basis, the four Re/TiO2
   states of the workbook (Candidate_Inputs, Table 3 inputs) must reproduce the stored workbook economics and the
   canonical ranking file to < 1e-6 EUR/t, and the 0.5-40 % purge sweep likewise (all 4 x 396 points).
2. Mass balance: C, H and O close for arbitrary inputs with CO recycle (x_CO > 0); x_CO = 0 equals the engine.
3. Reference loop (Processes 2022, 10, 1535): inert-CO vs recycled-CO treatments against the published loop
   (recycle 54,290 kmol/h, NPC 1071.8 MEUR/y, reactor CO 1.50 % in / 1.76 % out, 0.5 % CO selectivity, equilibrium
   CO2 conversion 30.4 %).
4. The four Re states with CO recycled (central rule and range): costs, ranking, purge optimum.

Writes validation_summary.json, reference_loop_comparison.csv, canonical_states_co_treatments.csv here and
exits non-zero if a reproduction gate fails.
"""
import csv
import json
import sys
from itertools import combinations
from pathlib import Path

import numpy as np
import openpyxl
from scipy.stats import kendalltau, spearmanr

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
sys.path.insert(0, str(REPO / "data" / "meoh"))
import meoh_d01_model as M  # noqa: E402
import meoh_general_model as G  # noqa: E402

WB = REPO / "data/meoh/MeOH_D01_ExplicitRecycleSeparationEconomics_v3.0.xlsx"
TOL = 1e-6
gates = {}

# ---------------------------------------------------------------- 1. canonical reproduction --------------------
wb = openpyxl.load_workbook(WB, data_only=True)
ci = wb["Candidate_Inputs"]
cands = []
for r in range(5, 9):
    name, re_wt, t_c, sty, x, smeoh, sch4, sco = (ci.cell(r, j).value for j in range(1, 9))
    cands.append(dict(name=name, metal_wt=float(re_wt), T_C=float(t_c), STY_per_g_metal=float(sty), X=float(x),
                      SMeOH=float(smeoh), SCH4=float(sch4), SCO=float(sco), P_bar=100.0, h2_co2=4.0))
names = [c["name"] for c in cands]

el = wb["Explicit_Loop_2pct"]
cols = ["elec_eur_t", "catalyst_t", "EC_MEUR", "sep_loop_EC_MEUR", "FCI_MEUR", "ACC_MEUR_Y", "direct_MEUR_Y",
        "indirect_MEUR_Y", "NPC_MEUR_Y"]
d_wb, d_wb_cols, d_engine = 0.0, 0.0, 0.0
npc_inert = {}
for i, c in enumerate(cands):
    g = G.cost(c, x_co="inert")
    e = M.candidate_economics(c["X"], c["SMeOH"], c["SCH4"], c["SCO"], c["STY_per_g_metal"], c["metal_wt"])
    npc_inert[c["name"]] = float(g["cost_eur_t"])
    for j, k in enumerate(cols):
        d_wb_cols = max(d_wb_cols, abs(float(g[k]) - float(el.cell(5 + i, 25 + j).value)))
    for k in list(cols) + ["cost_eur_t", "recycle_kmol_h", "h2_eur_t", "nonreactive_fraction"]:
        d_engine = max(d_engine, abs(float(g[k]) - float(e[k])))
    d_wb = max(d_wb, abs(float(g["NPC_MEUR_Y"]) - float(el.cell(5 + i, 33).value)) * 1e6 / M.PROD_TPY)

rank_csv = {r["candidate"]: r for r in csv.DictReader(open(REPO / "data/meoh/meoh_candidate_ranking_D01v3.csv",
                                                           encoding="utf-8"))}
d_csv_rounded = max(abs(round(npc_inert[n], 2) - float(rank_csv[n]["NPC_EUR_t_2pct_purge"])) for n in names)
d_csv_raw = max(abs(npc_inert[n] - float(rank_csv[n]["NPC_EUR_t_2pct_purge"])) for n in names)
econ_rank = {n: i + 1 for i, n in enumerate(sorted(names, key=lambda n: npc_inert[n]))}
rank_ok = all(econ_rank[n] == int(rank_csv[n]["economic_rank"]) for n in names)
gates["canonical_NPC_vs_workbook_max_abs_EUR_t"] = d_wb
gates["canonical_stored_columns_Y_to_AG_vs_workbook_max_abs"] = d_wb_cols
gates["canonical_all_columns_vs_engine_max_abs"] = d_engine
gates["canonical_NPC_vs_ranking_csv_after_2dp_rounding_max_abs"] = d_csv_rounded
gates["canonical_NPC_vs_ranking_csv_raw_max_abs (csv stores 2 dp)"] = d_csv_raw
gates["canonical_economic_rank_matches_csv"] = rank_ok

# purge sweep vs workbook Purge_Sweep (stored values) and the purge robustness file
ps = wb["Purge_Sweep"]
sweep_cols = ["cost_eur_t", "recycle_kmol_h", "nonreactive_fraction", "methane_fraction", "h2_eur_t", "elec_eur_t",
              "EC_MEUR"]
sweeps = {c["name"]: G.purge_sweep(c, x_co="inert") for c in cands}
d_sweep_cost, d_sweep_all, nrows, r = 0.0, {}, 0, 5
for c in cands:
    s = sweeps[c["name"]]
    for k, p in enumerate(G.PURGES):
        assert ps.cell(r, 1).value == c["name"] and abs(ps.cell(r, 2).value - p) < 1e-12
        for j, col in enumerate(sweep_cols):
            dv = abs(float(s[col][k]) - float(ps.cell(r, 3 + j).value))
            d_sweep_all[col] = max(d_sweep_all.get(col, 0.0), dv)
        r += 1
        nrows += 1
d_sweep_cost = d_sweep_all["cost_eur_t"]
rob = list(csv.DictReader(open(REPO / "data/meoh/meoh_purge_robustness_D01v3.csv", encoding="utf-8")))
sty = {c["name"]: c["STY_per_g_metal"] for c in cands}
up = {n: i + 1 for i, n in enumerate(sorted(names, key=lambda n: -sty[n]))}
d_rob, rob_rank_ok = 0.0, True
for row in rob:
    k = int(round((float(row["purge"]) - 0.005) / 0.001))
    npc = {n: float(sweeps[n]["cost_eur_t"][k]) for n in names}
    d_rob = max(d_rob, max(abs(round(npc[n], 2) - float(row["NPC_" + n])) for n in names))
    order = sorted(names, key=lambda n: npc[n])
    er = {n: i + 1 for i, n in enumerate(order)}
    a, b = [up[n] for n in names], [er[n] for n in names]
    inv = sum(1 for x, y in combinations(names, 2) if (up[x] - up[y]) * (er[x] - er[y]) < 0)
    rob_rank_ok &= (" > ".join(order) == row["economic_order"] and inv == int(row["pairwise_inversions"])
                    and round(float(spearmanr(a, b).statistic), 3) == float(row["spearman"]))
gates["purge_sweep_points"] = nrows
gates["purge_sweep_NPC_vs_workbook_max_abs_EUR_t"] = d_sweep_cost
gates["purge_sweep_all_columns_vs_workbook_max_abs"] = d_sweep_all
gates["purge_sweep_vs_robustness_csv_after_2dp_rounding_max_abs"] = d_rob
gates["purge_sweep_orders_inversions_spearman_match_robustness_csv"] = bool(rob_rank_ok)

# ---------------------------------------------------------------- 2. mass balance / engine identity -----------
rng = np.random.default_rng(20261005)
n = 2000
X = rng.uniform(0.03, 0.6, n)
smeoh = rng.uniform(0.3, 1.0, n)
rest = 1 - smeoh
fco = rng.uniform(0, 1, n)
sco, sch4 = rest * fco, rest * (1 - fco)
ratio = rng.uniform(2.0, 6.0, n)
purge = rng.uniform(0.005, 0.4, n)
xco = rng.uniform(0, 1, n)
lb = G.loop_balance(X, smeoh, sch4, sco, ratio, purge, xco)
p = purge
c_out = 1 + p * (lb["co2_out"] + lb["co_out"] + lb["ch4_in"] + lb["ch4_prod"])
h_out = 4 + 2 * lb["water_prod"] + p * (2 * lb["h2_out"] + 4 * (lb["ch4_in"] + lb["ch4_prod"]))
o_out = 1 + lb["water_prod"] + p * (2 * lb["co2_out"] + lb["co_out"])
bal = dict(C=float(np.max(np.abs(c_out - lb["fresh_co2"]))),
           H=float(np.max(np.abs(h_out - 2 * lb["fresh_h2"] * 1.0))),
           O=float(np.max(np.abs(o_out - 2 * lb["fresh_co2"]))))
lb0 = G.loop_balance(X, smeoh, sch4, sco, ratio, purge, 0.0)
le = M.loop_balance(X, smeoh, sch4, sco, ratio, purge)
d_loop = max(float(np.max(np.abs(lb0[k] - le[k]))) for k in le)
gates["mass_balance_max_abs_error_per_mol_MeOH (2000 random states with CO recycle)"] = bal
gates["x_co_0_loop_vs_engine_max_abs (2000 random states)"] = d_loop

ok = (d_wb < TOL and d_wb_cols < TOL and d_engine < TOL and d_csv_rounded < 1e-9 and rank_ok and d_sweep_cost < TOL
      and max(d_sweep_all.values()) < TOL and d_rob < 1e-9 and rob_rank_ok and max(bal.values()) < 1e-9
      and d_loop < 1e-9)
gates["ALL_REPRODUCTION_GATES_PASS"] = bool(ok)

# ---------------------------------------------------------------- 3. reference loop ---------------------------
SRC = dict(recycle_kmol_h=54290.0, NPC_MEUR_Y=1071.8, co_in_pct=1.50, co_out_pct=1.76, n2_in_pct=4.95,
           pass_net_S_CO=0.005, meoh_out_pct=7.4, h2o_out_pct=7.2, X_eq=0.304)
# published reactor outlet reconstructed from the inlet (Figure 6) and the outlet CO, N2, MeOH, H2O (section 3.1.2)
f_in = G.REF_FEED_MOLPCT
tot_out = f_in["N2"] / 0.0565
out = dict(MeOH=0.074 * tot_out, H2O=0.072 * tot_out, CO=0.0176 * tot_out, N2=f_in["N2"], CH4=0.0)
out["CO2"] = f_in["CO2"] + f_in["CO"] + f_in["CH3OH"] - out["CO"] - out["MeOH"]
out["H2"] = tot_out - sum(out.values())
y = {k: v / tot_out for k, v in out.items()}
T_ref = G.REF_T_C + 273.15
fa = G.gas_activities(y, T_ref, M.REF_PRESSURE_BAR)
src_rwgs = float(fa["CO"] * fa["H2O"] / (fa["CO2"] * fa["H2"]) / G.K_rwgs(T_ref))
src_co2h = float(fa["MeOH"] * fa["H2O"] / (fa["CO2"] * fa["H2"] ** 3) / G.K_co2_hyd(T_ref))
src_X = (f_in["CO2"] - out["CO2"]) / f_in["CO2"]
feed = dict(H2=f_in["H2"], CO=f_in["CO"], CO2=f_in["CO2"], MeOH=f_in["CH3OH"], H2O=0.0, N2=f_in["N2"])
X_eq_pr, _ = G.equilibrium_co2_conversion(feed, G.REF_T_C, M.REF_PRESSURE_BAR, "PR")
X_eq_id, _ = G.equilibrium_co2_conversion(feed, G.REF_T_C, M.REF_PRESSURE_BAR, "ideal")

ref_rows = []
cases = [("inert (x_CO = 0)", "inert"), ("recycled, central (x_RWGS)", "recycled_central"),
         ("recycled, high (thermodynamic max)", "recycled_high"),
         ("recycled, equal-X sensitivity", "recycled_equal_X"),
         ("recycled at x_CO = X = 0.285, unclipped (0.5 % read as gross)", M.SOURCE_X)]
for label, rule in cases:
    e = G.reference_economics(rule)
    lp = e["loop"]
    ref_rows.append({
        "case": label, "x_CO": float(e["x_co"]), "NPC_MEUR_y": float(e["NPC_MEUR_Y"]),
        "NPC_EUR_t": float(e["cost_eur_t"]), "recycle_kmol_h": float(e["recycle_kmol_h"]),
        "raw_recycle_kmol_h": float(e["raw_recycle_kmol_h"]),
        "reactor_in_CO_pct": 100 * float(e["co_inlet_fraction"]), "reactor_out_CO_pct": 100 * float(e["y_out"]["CO"]),
        "reactor_in_N2_pct": 100 * float(e["n2_inlet_fraction"]),
        "reactor_in_H2_pct": 100 * float(lp["h2_in"] / lp["reactor_in"]),
        "reactor_in_CO2_pct": 100 * float(lp["co2_in"] / lp["reactor_in"]),
        "reactor_out_MeOH_pct": 100 * float(e["y_out"]["MeOH"]), "reactor_out_H2O_pct": 100 * float(e["y_out"]["H2O"]),
        "pass_net_S_CO": float(e["pass_net_co_selectivity"]), "overall_net_S_CO": float(e["net_co_selectivity"]),
        "RWGS_Q_over_K_outlet": float(e["rwgs_approach"]), "CO2hyd_Q_over_K_outlet": float(e["co2_hyd_approach"]),
        "carbon_efficiency": float(e["carbon_efficiency"])})
ref_rows.append({"case": "published loop (Processes 2022, 10, 1535)", "x_CO": None, "NPC_MEUR_y": SRC["NPC_MEUR_Y"],
                 "NPC_EUR_t": 920.0, "recycle_kmol_h": SRC["recycle_kmol_h"], "raw_recycle_kmol_h": None,
                 "reactor_in_CO_pct": 1.50, "reactor_out_CO_pct": 1.76, "reactor_in_N2_pct": 5.0,
                 "reactor_in_H2_pct": 71.3, "reactor_in_CO2_pct": 21.9, "reactor_out_MeOH_pct": 7.4,
                 "reactor_out_H2O_pct": 7.2, "pass_net_S_CO": 0.005, "overall_net_S_CO": None,
                 "RWGS_Q_over_K_outlet": src_rwgs, "CO2hyd_Q_over_K_outlet": src_co2h, "carbon_efficiency": 0.943})
with open(HERE / "reference_loop_comparison.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(ref_rows[0]))
    w.writeheader()
    w.writerows(ref_rows)

# ---------------------------------------------------------------- 4. four Re states, CO treatments -------------
treat = ["inert", "recycled_central", "recycled_high", "recycled_equal_X"]
rows4, res4 = [], {}
for c in cands:
    for t in treat:
        e = G.cost(c, x_co=t)
        s = G.purge_sweep(c, x_co=t)
        k = int(np.argmin(s["cost_eur_t"]))
        res4[(c["name"], t)] = dict(npc=float(e["cost_eur_t"]), opt=float(s["cost_eur_t"][k]),
                                    opt_purge=float(G.PURGES[k]))
        rows4.append({"candidate": c["name"], "treatment": t, "x_CO_2pct": float(e["x_co"]),
                      "NPC_EUR_t_2pct": float(e["cost_eur_t"]), "delta_vs_inert_EUR_t":
                      float(e["cost_eur_t"]) - npc_inert[c["name"]],
                      "recycle_kmol_h": float(e["recycle_kmol_h"]),
                      "reactor_in_CO_pct": 100 * float(e["co_inlet_fraction"]),
                      "nonH2CO2_fraction": float(e["nonreactive_fraction"]),
                      "RWGS_Q_over_K_outlet": float(e["rwgs_approach"]),
                      "opt_purge": float(G.PURGES[k]), "NPC_EUR_t_opt": float(s["cost_eur_t"][k])})
with open(HERE / "canonical_states_co_treatments.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(rows4[0]))
    w.writeheader()
    w.writerows(rows4)


def rank_metrics(cost):
    order = sorted(names, key=lambda n: cost[n])
    er = {n: i + 1 for i, n in enumerate(order)}
    a, b = [up[n] for n in names], [er[n] for n in names]
    inv = sum(1 for x, y in combinations(names, 2) if (up[x] - up[y]) * (er[x] - er[y]) < 0)
    upw = min(names, key=lambda n: up[n])
    return dict(order=" > ".join(order), spearman=float(spearmanr(a, b).statistic),
                kendall=float(kendalltau(a, b).statistic), inversions=inv,
                upstream_winner_econ_rank=er[upw],
                regret=(cost[upw] - cost[order[0]]) / cost[order[0]])


four = {}
for t in treat:
    for basis in ("npc", "opt"):
        four["%s | %s" % (t, "2 % purge" if basis == "npc" else "optimum purge")] = rank_metrics(
            {n: res4[(n, t)][basis] for n in names})

summary = dict(
    gates=gates,
    reference_loop=dict(
        published=SRC, published_outlet_reconstructed=dict(CO2_conversion=src_X, RWGS_Q_over_K=src_rwgs,
                                                          CO2hyd_Q_over_K=src_co2h),
        equilibrium_CO2_conversion_of_published_feed=dict(PR=X_eq_pr, ideal_gas=X_eq_id, published=0.304),
        model=ref_rows[:-1]),
    canonical_states=dict(costs={"%s | %s" % k: v for k, v in res4.items()}, ranking_vs_STY_per_gRe=four),
)
(HERE / "validation_summary.json").write_text(json.dumps(summary, indent=1), encoding="utf-8")

print(json.dumps(gates, indent=1))
print("published outlet: X %.4f  RWGS Q/K %.4f  CO2-hyd Q/K %.4f ; X_eq PR %.4f ideal %.4f (published 0.304)"
      % (src_X, src_rwgs, src_co2h, X_eq_pr, X_eq_id))
for r_ in ref_rows:
    print("%-62s x %-8s NPC %8.3f  rec %8.0f  CO in %.2f out %.2f  RWGS %.3f  Snet %.4f" % (
        r_["case"], "-" if r_["x_CO"] is None else "%.5f" % r_["x_CO"], r_["NPC_MEUR_y"], r_["recycle_kmol_h"],
        r_["reactor_in_CO_pct"], r_["reactor_out_CO_pct"], r_["RWGS_Q_over_K_outlet"], r_["pass_net_S_CO"]))
for r_ in rows4:
    print("%-18s %-18s x %.4f  NPC %8.3f (%+.3f)  opt %8.3f @ %.3f  RWGS %.3f" % (
        r_["candidate"], r_["treatment"], r_["x_CO_2pct"], r_["NPC_EUR_t_2pct"], r_["delta_vs_inert_EUR_t"],
        r_["NPC_EUR_t_opt"], r_["opt_purge"], r_["RWGS_Q_over_K_outlet"]))
for k, v in four.items():
    print("%-36s %s  rho %.2f tau %.2f inv %d  STY-winner #%d regret %.4f%%" % (
        k, v["order"], v["spearman"], v["kendall"], v["inversions"], v["upstream_winner_econ_rank"],
        100 * v["regret"]))
if not ok:
    sys.exit("REPRODUCTION GATE FAILED")
