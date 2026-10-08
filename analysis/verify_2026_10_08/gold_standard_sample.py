"""Draw the human-checked gold-standard sample for the methanol extraction (protocol fixed before the draw).

Protocol
--------
Population: every value that enters the methanol field result on printed values only, i.e. the entries of the
printed-value candidate set (meoh_decomposition.build(..., printed_only=True), corrected set S3) that sit in a scored
comparison group (at least two entries). For each such entry the values that enter the calculation are:
CO2 conversion, methanol selectivity, CO and CH4 selectivity when reported, temperature, pressure, H2/CO2, the
space velocity when the group's productivity is derived from it, and the printed STY when the group uses it.

Draw: 10 papers uniformly at random without replacement from the papers in the population (numpy default_rng,
SEED); then 10 values per paper uniformly at random without replacement from that paper's values. A paper with
fewer than 10 values contributes all of them, and the shortfall is drawn uniformly from the remaining values of the
other drawn papers, so the sample has exactly 100 values.

Check (by hand, against the paper or its Supporting Information, not by any program or model): for every value the
checker records where it is printed, the value printed there, whether it is a printed number (not read off a plot),
and the judgment: correct (equal to the printed value after unit conversion), rounding (differs only beyond the
printed digits), or wrong. Accuracy = correct / 100 with a Wilson 95 % interval; rounding is reported separately.

Output: gold_standard_sample.csv (this folder) and the workbook for the checker (path given as the first argument).
"""
import sys
from io import StringIO
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import meoh_decomposition as D  # noqa: E402

SEED = 20261008
N_PAPERS, PER_PAPER = 10, 10
MANIFESTS = [Path(r"D:\论文-AI4S\wt-extraction40\agent\extraction"), Path(r"D:\论文-AI4S\Catalyst-Essay\agent\extraction")]


def population():
    rec = pd.read_csv(StringIO(D.git_show("agent/extraction/out/records_normalized.csv")), keep_default_na=False,
                      na_values=[""])
    rec["rid"] = np.arange(len(rec))
    d = rec.dropna(subset=D.MC.REQUIRED)
    rows = []
    for r in d.itertuples(index=False):
        c = D.candidate(r, True)
        if c is not None:
            c["rid"] = r.rid
            rows.append(c)
    cand = pd.DataFrame(rows)
    cand["group"] = (cand.doi + " | " + cand.P_bar.map("{:g} bar".format) + " | H2/CO2 "
                     + cand.h2_co2.map("{:.3g}".format) + " | " + cand.ghsv_key)
    parts = []
    for _, g in cand.groupby("group"):
        b, sty = D.basis_lock(g)
        g = g.copy()
        g["sty_basis"], g["STY"] = b, sty
        parts.append(g)
    cand = pd.concat([p for p in parts if p.STY.notna().any()]).dropna(subset=["STY"])
    cand = cand[cand.STY > 0]
    cand = cand[~cand.duplicated(["group", "T_C", "X", "SMeOH", "SCH4", "SCO", "STY"], keep="first")]
    n = cand.groupby("group").size()
    cand = cand[cand.group.isin(n[n >= 2].index)]
    # one row per value that enters the calculation
    vals = []
    for c in cand.itertuples(index=False):
        r = rec.iloc[int(c.rid)]
        def add(field, value, unit, src=None, q=None):
            vals.append(dict(doi=c.doi, entry=c.entry, catalyst=r.catalyst_name, group=c.group, field=field,
                             extracted=value, unit=unit, qualifier=q, value_source=src,
                             record_location=r.primary_location, record_page=r.primary_page, pass_=r["pass"],
                             filled=r.filled if isinstance(r.filled, str) else ""))
        add("CO2 conversion", r.X_CO2_pct, "%", r.X_CO2_src, r.X_CO2_q)
        add("CH3OH selectivity", r.S_MeOH_pct, "%", r.S_MeOH_src, r.S_MeOH_q)
        if pd.notna(r.S_CO_pct):
            add("CO selectivity", r.S_CO_pct, "%", r.S_CO_src, r.S_CO_q)
        if pd.notna(r.S_CH4_pct):
            add("CH4 selectivity", r.S_CH4_pct, "%", r.S_CH4_src, r.S_CH4_q)
        add("temperature", round(float(r.T_K) - 273.15, 2), "°C (from K)")
        add("pressure", r.P_bar, "bar")
        add("H2/CO2", r.H2_CO2, "mol/mol")
        if c.sty_basis.startswith("printed STY"):
            add("STY (printed)", r.STY_raw, "as printed", r.STY_src, r.STY_q)
        else:
            add("space velocity", r.GHSV_raw, "as printed", r.GHSV_src)
    return pd.DataFrame(vals), cand


def pdf_index():
    import json
    idx = {}
    for base in MANIFESTS:
        for m in json.loads((base / "fetch_manifest.json").read_text(encoding="utf-8")):
            p = base / "pdf" / str(m.get("file") or "")
            if m.get("file") and p.exists():
                idx.setdefault(m["doi"], str(p))
        si = base / "si_manifest.json"
        if si.exists():
            for m in json.loads(si.read_text(encoding="utf-8")):
                names = [f["file"] if isinstance(f, dict) else f for f in (m.get("files") or [])]
                files = [str(base / "si" / f) for f in names if (base / "si" / f).exists()]
                if files:
                    idx.setdefault(m["doi"] + "#si", "; ".join(files))
    return idx


def main(out_xlsx):
    vals, cand = population()
    rng = np.random.default_rng(SEED)
    papers = np.array(sorted(vals.doi.unique()))
    drawn = list(rng.choice(papers, size=N_PAPERS, replace=False))
    picks, rest = [], []
    for doi in drawn:
        pool = vals[vals.doi == doi]
        k = min(PER_PAPER, len(pool))
        take = rng.choice(pool.index.to_numpy(), size=k, replace=False)
        picks += list(take)
        rest += [i for i in pool.index if i not in set(take)]
    short = N_PAPERS * PER_PAPER - len(picks)
    if short > 0:
        picks += list(rng.choice(np.array(rest), size=short, replace=False))
    s = vals.loc[picks].copy()
    s.insert(0, "id", np.arange(1, len(s) + 1))
    idx = pdf_index()
    s["pdf"] = s.doi.map(lambda d: idx.get(d, ""))
    s["si"] = s.doi.map(lambda d: idx.get(d + "#si", ""))
    for col in ("found_where", "value_in_paper", "printed_number", "judgment", "note"):
        s[col] = ""
    s.to_csv(HERE / "gold_standard_sample.csv", index=False, encoding="utf-8-sig")
    meta = dict(seed=SEED, population_values=len(vals), population_entries=len(cand),
                population_papers=int(vals.doi.nunique()), drawn_papers=drawn, shortfall_topup=max(short, 0))
    print(meta)
    write_xlsx(s, meta, out_xlsx)


def write_xlsx(s, meta, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    lines = [
        "甲醇文献抽取：人工金标准核对表",
        "",
        "抽样规则（抽样前固定，见 analysis/verify_2026_10_08/gold_standard_sample.py）：",
        f"总体 = 只用印刷值时进入甲醇主结果的全部数值：{meta['population_values']} 个数，"
        f"{meta['population_entries']} 个条目，{meta['population_papers']} 篇论文。",
        f"随机种子 {meta['seed']}：先随机抽 10 篇论文，每篇随机抽 10 个数，共 100 个数。",
        "",
        "核对方法（必须本人对照原文或 SI 完成，不用任何程序或模型判断）：",
        "1. 打开 pdf / si 列给出的文件，按 entry、record_location、record_page 找到该条目。",
        "2. found_where：写出这个数在原文印刷的位置（如 'Table 2, p.5' 或 'SI Table S3'）。",
        "3. value_in_paper：抄下原文印刷的数（原单位）。",
        "4. printed_number：原文是印刷的数字填 Y；只能从图上读填 N。",
        "5. judgment：correct = 换算单位后与原文一致；rounding = 只在原文印刷位数之外有差；wrong = 不一致。",
        "6. note：其他说明（如单位换算、原文自相矛盾）。",
        "",
        "字段说明：temperature 已从 K 换算成 °C；pressure 单位 bar（1 MPa = 10 bar）；",
        "STY (printed) 和 space velocity 给的是原文写法（as printed）。",
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).font = Font(size=12, bold=(i == 1))
    ws.column_dimensions["A"].width = 120
    sh = wb.create_sheet("核对表")
    cols = ["id", "doi", "catalyst", "entry", "field", "extracted", "unit", "record_location", "record_page",
            "value_source", "filled", "pdf", "si", "found_where", "value_in_paper", "printed_number", "judgment", "note"]
    sh.append(cols)
    for r in s[cols].itertuples(index=False):
        sh.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r])
    fill = PatternFill("solid", fgColor="FFF2CC")
    for j, c in enumerate(cols, 1):
        sh.cell(row=1, column=j).font = Font(bold=True)
        if c in ("found_where", "value_in_paper", "printed_number", "judgment", "note"):
            for i in range(1, len(s) + 2):
                sh.cell(row=i, column=j).fill = fill
    widths = dict(id=5, doi=26, catalyst=28, entry=40, field=18, extracted=12, unit=12, record_location=28,
                  record_page=8, value_source=10, filled=18, pdf=30, si=30, found_where=22, value_in_paper=14,
                  printed_number=10, judgment=12, note=30)
    for j, c in enumerate(cols, 1):
        sh.column_dimensions[sh.cell(row=1, column=j).column_letter].width = widths[c]
    for row in sh.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    n = len(s) + 1
    dv1 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"correct,rounding,wrong"', allow_blank=True)
    sh.add_data_validation(dv1)
    sh.add_data_validation(dv2)
    dv1.add(f"P2:P{n}")
    dv2.add(f"Q2:Q{n}")
    sh.freeze_panes = "B2"
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(HERE / "gold_standard_sample.xlsx"))
