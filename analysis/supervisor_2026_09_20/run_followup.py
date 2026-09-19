from __future__ import annotations
import csv, hashlib, json, math, sys
from pathlib import Path
import numpy as np
import openpyxl, yaml
from scipy.optimize import brentq

ROOT=Path(__file__).resolve().parents[2]
SRC=ROOT/"provenance/discover_v1/source_harness"
sys.path.insert(0,str(SRC))
import harness_core as HC
from discover.env import DiscoverEnv

OUT=ROOT/"data/supervisor_2026_09_20"
OUT.mkdir(parents=True,exist_ok=True)
N=1000
SEED=20260920

def sha256(p:Path):
    h=hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda:f.read(1<<20),b""): h.update(b)
    return h.hexdigest()

cfg_path=SRC/"configs/nh3_final.yaml"
wb_path=SRC/"ammonia_activity_volcano_s1_v1.xlsx"
cfg=yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
h=HC.NH3Harness(cfg,SRC)
fe_vec=h.frozen_state_logtof(h.EN0["Fe"])
ru_vec=h.frozen_state_logtof(h.EN0["Ru"])

def evaluate(metal,vec,*,price=None,life_y=None,alpha=1.0,capex_multiplier=1.0,electricity=None):
    if price is None: price=HC.PRICE[metal]
    if life_y is None: life_y=h.CATALYST_LIFE_Y
    if electricity is None: electricity=h.ELECTRICITY
    _,V,metal_cost,reactor=h.cost_arrays(metal,vec,price=price,life_y=life_y,alpha=alpha)
    elec_scale=float(electricity)/h.ELECTRICITY
    opex=(h.state_fresh+h.state_reccomp+h.state_refrig)*elec_scale
    compcap=h.state_compcapex
    total=metal_cost+capex_multiplier*(reactor+compcap)+opex
    ok=V<=h.V_CAP
    i=int(np.argmin(np.where(ok,total,np.inf)))
    vi=float(V[i])
    reactor_base=float((h.REACTOR_FIXED+h.REACTOR_VAR*(vi/h.REACTOR_REF)**h.REACTOR_EXP)*h.crf/h.annual_output_t)
    bd={
        "metal_cost":float(metal_cost[i]),
        "reactor_base":reactor_base*capex_multiplier,
        "vessel_pressure_premium":(float(reactor[i])-reactor_base)*capex_multiplier,
        "fresh_comp":float(h.state_fresh[i])*elec_scale,
        "recycle_comp":float(h.state_reccomp[i])*elec_scale,
        "refrigeration":float(h.state_refrig[i])*elec_scale,
        "compressor_capex":float(compcap[i])*capex_multiplier,
    }
    return {"total_cost":float(total[i]),"V_m3":vi,"T_C":float(h.state_T[i]),"P_bar":float(h.state_P[i]),"Tsep_C":float(h.state_Tsep[i]),"state_index":i,"breakdown":bd}

# 1) Ru price-equalization counterfactual
fe=evaluate("Fe",fe_vec)
ru=evaluate("Ru",ru_vec)
ru_equal=evaluate("Ru",ru_vec,price=HC.PRICE["Fe"])
cf={
  "provenance":{"config_sha256":sha256(cfg_path),"activity_workbook_sha256":sha256(wb_path),"source_harness":"provenance/discover_v1/source_harness"},
  "baseline":{"Fe":fe,"Ru":ru,"gap_Ru_minus_Fe":ru["total_cost"]-fe["total_cost"]},
  "counterfactual_Ru_price_equals_Fe":{"Ru":ru_equal,"Fe_reference":fe,"gap_Ru_minus_Fe":ru_equal["total_cost"]-fe["total_cost"],"winner":"Ru" if ru_equal["total_cost"]<fe["total_cost"] else "Fe","price_Ru_baseline_USD_kg":HC.PRICE["Ru"],"price_Ru_counterfactual_USD_kg":HC.PRICE["Fe"]}
}
(OUT/"nh3_ru_price_counterfactual.json").write_text(json.dumps(cf,indent=2),encoding="utf-8")
components=list(fe["breakdown"])
with (OUT/"nh3_gap_decomposition.csv").open("w",newline="",encoding="utf-8") as f:
    w=csv.writer(f); w.writerow(["component","baseline_Ru_minus_Fe_USD_t","equal_price_Ru_minus_Fe_USD_t"])
    for k in components: w.writerow([k,ru["breakdown"][k]-fe["breakdown"][k],ru_equal["breakdown"][k]-fe["breakdown"][k]])
    w.writerow(["TOTAL",ru["total_cost"]-fe["total_cost"],ru_equal["total_cost"]-fe["total_cost"]])

# 2) Frozen joint cost-side MC for NH3
rng=np.random.default_rng(SEED)
rows=[]
for d in range(N):
    fe_pm=float(rng.uniform(.5,1.5)); ru_pm=float(rng.uniform(.5,1.5))
    cap=float(rng.uniform(.8,1.2)); elec=float(rng.uniform(20.,100.)); life=float(rng.uniform(5.,20.))
    fe_d=evaluate("Fe",fe_vec,price=HC.PRICE["Fe"]*fe_pm,life_y=life,capex_multiplier=cap,electricity=elec)
    ru_d=evaluate("Ru",ru_vec,price=HC.PRICE["Ru"]*ru_pm,life_y=life,capex_multiplier=cap,electricity=elec)
    target=fe_d["total_cost"]
    def gap(loga):
        return evaluate("Ru",ru_vec,price=HC.PRICE["Ru"]*ru_pm,life_y=life,alpha=10.**loga,capex_multiplier=cap,electricity=elec)["total_cost"]-target
    if gap(0.)<=0: alpha=1.0
    elif gap(12.)>0: alpha=float("nan")
    else: alpha=float(10.**brentq(gap,0.,12.,xtol=1e-10,rtol=1e-10,maxiter=100))
    rows.append({"draw":d,"Fe_price_multiplier":fe_pm,"Ru_price_multiplier":ru_pm,"capex_multiplier":cap,"electricity_USD_MWh":elec,"catalyst_life_y":life,
                 "Fe_cost_USD_t":fe_d["total_cost"],"Ru_cost_USD_t":ru_d["total_cost"],"gap_Ru_minus_Fe_USD_t":ru_d["total_cost"]-fe_d["total_cost"],
                 "Fe_lower_than_Ru":int(fe_d["total_cost"]<ru_d["total_cost"]),"alpha_star":alpha,
                 "Fe_P_bar":fe_d["P_bar"],"Ru_P_bar":ru_d["P_bar"]})
with (OUT/"nh3_joint_cost_mc_draws.csv").open("w",newline="",encoding="utf-8") as f:
    w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
arr=lambda k:np.array([r[k] for r in rows],float)
a=arr("alpha_star"); gapv=arr("gap_Ru_minus_Fe_USD_t")
nh3_mc={
  "protocol":"docs/JOINT_COST_MC_PROTOCOL_2026-09-20.md","draws":N,"seed":SEED,
  "P_C_Fe_lt_C_Ru":float(np.mean(arr("Fe_lower_than_Ru"))),"Fe_wins_count":int(np.sum(arr("Fe_lower_than_Ru"))),
  "cost_gap_Ru_minus_Fe_USD_t":{"p05":float(np.quantile(gapv,.05)),"median":float(np.median(gapv)),"p95":float(np.quantile(gapv,.95)),"min":float(np.min(gapv)),"max":float(np.max(gapv))},
  "alpha_star":{"finite_count":int(np.isfinite(a).sum()),"p05":float(np.nanquantile(a,.05)),"median":float(np.nanmedian(a)),"p95":float(np.nanquantile(a,.95)),"min":float(np.nanmin(a)),"max":float(np.nanmax(a))}
}
(OUT/"nh3_joint_cost_mc_summary.json").write_text(json.dumps(nh3_mc,indent=2),encoding="utf-8")

# 3) F3 metric reconciliation
mc_detail_path=SRC/"outputs/nh3_final_1_1_20260905T113727Z/closure/mc_detail.json"
if not mc_detail_path.exists():
    mc_detail_path=ROOT/"provenance/nh3_final_1_1/source_harness/outputs/nh3_final_20260905T134204Z/closure/mc_detail.json"
md=json.loads(mc_detail_path.read_text(encoding="utf-8"))
wt=md["winner_transition_fraction"]
fe_winner=sum(v for k,v in wt.items() if k.endswith("->Fe"))
f3={"source":str(mc_detail_path.relative_to(ROOT)),"Fe_final_economic_winner_probability":fe_winner,
    "top1_survival_metric":md["summary"]["top1_survival"],"top1_raw_before_feasibility":md["top1_raw_before_feasibility"],
    "Fe_feasibility_probability":md["summary"]["Fe_feasibility_probability"],"top3_actionable":md["summary"]["top3_actionable"],
    "winner_transition_fraction":wt,
    "interpretation":"The supervisor's ~68% Fe-first statement corresponds to the final economic-winner frequency (sum of transitions ending in Fe). The 28.2% field is top1_survival/top1_raw_before_feasibility and is not Fe's final economic-winner probability."}
(OUT/"f3_metric_reconciliation.json").write_text(json.dumps(f3,indent=2),encoding="utf-8")

# 4) MeOH cost-side rank probability from exact workbook economics
meoh_wb=ROOT/"data/meoh/MeOH_D01_ExplicitRecycleSeparationEconomics_v3.0.xlsx"
wb=openpyxl.load_workbook(meoh_wb,data_only=True,read_only=True)
ws=wb["Explicit_Loop_2pct"]; ci=wb["Candidate_Inputs"]; ctl=wb["Controls"]
annual_t=float(ctl["B5"].value)*float(ctl["B6"].value)
cands=[]
for row in range(5,9):
    cands.append({"name":ci.cell(row,1).value,"Re_wt_pct":float(ci.cell(row,2).value),"catalyst_t":float(ws.cell(row,26).value),
                  "FCI_MEUR":float(ws.cell(row,29).value),"ACC_MEUR_y":float(ws.cell(row,30).value),
                  "direct_MEUR_y":float(ws.cell(row,31).value),"indirect_MEUR_y":float(ws.cell(row,32).value),
                  "NPC_MEUR_y":float(ws.cell(row,33).value),"compression_EUR_t":float(ws.cell(row,25).value)})
rng=np.random.default_rng(SEED)
counts0={c["name"]:[0]*4 for c in cands}; counts_ext={c["name"]:[0]*4 for c in cands}
re_dist={c["name"]:[] for c in cands}
for d in range(N):
    re_pm=float(rng.uniform(.5,1.5)); cap=float(rng.uniform(.8,1.2)); elec=float(rng.uniform(20.,100.)); life=float(rng.uniform(5.,20.))
    base_cost=[]; ext_cost=[]
    for c in cands:
        labor_term=c["indirect_MEUR_y"]-.081*c["FCI_MEUR"]-.10*c["NPC_MEUR_y"]
        direct_new=c["direct_MEUR_y"]+c["compression_EUR_t"]*annual_t/1e6*(elec/90.-1.)
        fci_new=c["FCI_MEUR"]*cap; acc_new=c["ACC_MEUR_y"]*cap
        npc_new=(acc_new+direct_new+labor_term+.081*fci_new)/.90
        per_t=npc_new*1e6/annual_t
        re_mass_kg=c["catalyst_t"]*1000.*c["Re_wt_pct"]/100.
        re_repl=re_mass_kg*(HC.PRICE["Re"]/1.13)*re_pm/life/annual_t
        base_cost.append((c["name"],per_t)); ext_cost.append((c["name"],per_t+re_repl)); re_dist[c["name"]].append(re_repl)
    for arrc,counts in [(base_cost,counts0),(ext_cost,counts_ext)]:
        order=sorted(arrc,key=lambda x:x[1]); ranks={n:i for i,(n,_) in enumerate(order)}
        for n,_ in arrc: counts[n][ranks[n]]+=1
meoh={"draws":N,"seed":SEED,
      "canonical_boundary":"MEOH-D01-v3 intentionally excludes Re purchase price and does not model catalyst lifetime/replacement. Under the frozen boundary, metal-price and lifetime draws are structurally inactive; CAPEX and electricity are propagated through the existing equations.",
      "diagnostic_extension":"A separate supporting sensitivity adds active-Re replacement only: Re inventory = catalyst mass x Re wt%; Re baseline price = 5757.64 USD/kg from the frozen metal-price table, converted at 1.13 USD/EUR; zero recovery; price multiplier U(0.5,1.5), lifetime U(5,20 y). This does not silently redefine canonical D01-v3.",
      "rank_probability_canonical_boundary":{},"rank_probability_with_Re_replacement_extension":{},"Re_replacement_EUR_t_quantiles":{}}
for c in cands:
    n=c["name"]; meoh["rank_probability_canonical_boundary"][n]={f"rank_{i+1}":counts0[n][i]/N for i in range(4)}
    meoh["rank_probability_with_Re_replacement_extension"][n]={f"rank_{i+1}":counts_ext[n][i]/N for i in range(4)}
    q=np.asarray(re_dist[n]); meoh["Re_replacement_EUR_t_quantiles"][n]={"p05":float(np.quantile(q,.05)),"median":float(np.median(q)),"p95":float(np.quantile(q,.95))}
(OUT/"meoh_joint_cost_rank_probability.json").write_text(json.dumps(meoh,indent=2),encoding="utf-8")

# 5) Oracle minimum CU audit: shortest chain that satisfies the environment's S1-S3 stopping criteria.
cost_model=json.loads((SRC/"DISCOVER_COST_MODEL_V1.json").read_text(encoding="utf-8"))
env=DiscoverEnv(cost_model=cost_model,budget=100.,seed=0,anonymous=False)
activity=env.COMPUTE_ACTIVITY(env.candidates)
window=env.BUILD_PROCESS_WINDOW({"T_C":[425],"P_bar":[180,190],"Tsep_C":[30,30]},name="oracle")
opt={m:env.OPTIMIZE_PROCESS(m,"oracle") for m in ["Fe","Ru","Os"]}
bw=env.BACKWARD(["Ru","Fe"],property="activity",window="oracle")
reach=env.TEST_REACHABILITY("Ru",property="activity",scope="reference",required_multiplier=bw["multiplier"])
stop=env.stopping_status()
spent=env.budget0-env.budget
oracle={"definition":"Perfect-information action selection is allowed, but hidden numerical answers cannot be supplied as action arguments. The oracle must satisfy the environment's S1/S2/S3 stopping criteria. The three-state window contains Fe's exact economic optimum and the canonical Ru parity state; Ru and Os are explicitly resolved, while all remaining candidates are screened by the declared activity/price dominance rule.",
        "window":window,"optimized":opt,"backward":bw,"reachability":reach,"stopping_status":stop,"ledger":env.ledger,
        "oracle_minimum_CU":spent,"ratios_to_oracle":{"strong_budget_75":75/spent,"strong_75_median_decision_stable_52":52/spent,"fixed_policy_threshold_206":206/spent,"nonbinding_median_decision_stable_566":566/spent,"nonbinding_median_final_714":714/spent},
        "scorer_only_note":"A shorter scorer-targeting chain can be constructed if S1-S3 are ignored. It is intentionally not used as the scientific oracle baseline."}
assert stop["S1_winner_identified"] and stop["S2_no_unresolved_candidate"] and stop["S3_reachability_classified_if_needed"]
assert abs(spent-22.)<1e-9
(OUT/"agent_oracle_minimum_CU.json").write_text(json.dumps(oracle,indent=2),encoding="utf-8")

print(json.dumps({"counterfactual_gap":cf["counterfactual_Ru_price_equals_Fe"]["gap_Ru_minus_Fe"],"P_Fe_lt_Ru":nh3_mc["P_C_Fe_lt_C_Ru"],"alpha_median":nh3_mc["alpha_star"]["median"],"Fe_final_winner_probability":fe_winner,"oracle_CU":spent,"meoh_rank_probability":meoh["rank_probability_canonical_boundary"]},indent=2))
