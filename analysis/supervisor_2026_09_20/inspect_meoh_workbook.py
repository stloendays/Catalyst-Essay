from __future__ import annotations
import json
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
WB=ROOT/"data/meoh/MeOH_D01_ExplicitRecycleSeparationEconomics_v3.0.xlsx"
OUT=ROOT/"data/supervisor_2026_09_20"
OUT.mkdir(parents=True,exist_ok=True)

wbv=openpyxl.load_workbook(WB,data_only=True,read_only=True)
wbf=openpyxl.load_workbook(WB,data_only=False,read_only=True)
report={"workbook":str(WB.relative_to(ROOT)),"sheet_names":wbv.sheetnames,"sheets":{}}
for s in wbv.sheetnames:
    wsv=wbv[s]; wsf=wbf[s]
    info={"max_row":wsv.max_row,"max_column":wsv.max_column,"nonempty_preview":[]}
    count=0
    for row in wsv.iter_rows():
        vals=[c.value for c in row]
        if any(v is not None for v in vals):
            frec=[wsf.cell(c.row,c.column).value for c in row]
            info["nonempty_preview"].append({"row":row[0].row,"values":vals[:20],"formulas":frec[:20]})
            count+=1
            if count>=80: break
    report["sheets"][s]=info
(OUT/"meoh_workbook_inspection.json").write_text(json.dumps(report,indent=2,default=str),encoding="utf-8")
print(json.dumps({"sheets":report["sheet_names"],"out":"data/supervisor_2026_09_20/meoh_workbook_inspection.json"},indent=2))

# trigger: 2026-09-20
