"""Build the Supplementary Information Word file from docs/SUPPLEMENTARY_INFORMATION.md and docs/SI_TABLES.md.

Layout as in tools/build_manuscript_docx.py (Times New Roman 12 pt, double-spaced text, A4 with 2.5 cm margins, page
numbers), without line numbers. Headings ("### Supplementary Note N | title", "### Supplementary Table N | title")
stay headings; markdown tables become plain Word tables with a visible grid, a repeated header row and rows that do
not split across pages; a table too wide for a portrait page gets a landscape section. Bracketed working-reference
citations [n] are renumbered in order of first citation in the SI, and the reference list is built from the
"## Working references" section of docs/MANUSCRIPT_MAIN_TEXT.md.

    CatalystForge/.venv/python tools/build_si_docx.py OUT.docx [--si NOTES.md] [--tables TABLES.md]
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
ap = argparse.ArgumentParser()
ap.add_argument("out", nargs="?", default=str(ROOT / "Supplementary_Information.docx"))
ap.add_argument("--si", default=str(ROOT / "docs/SUPPLEMENTARY_INFORMATION.md"))
ap.add_argument("--tables", default=str(ROOT / "docs/SI_TABLES.md"))
args = ap.parse_args()

FONT = "Times New Roman"
NOTES = Path(args.si).read_text(encoding="utf-8")
TABLES = Path(args.tables).read_text(encoding="utf-8")
MAIN = (ROOT / "docs/MANUSCRIPT_MAIN_TEXT.md").read_text(encoding="utf-8")
REFS = {int(m.group(1)): m.group(2).strip()
        for m in re.finditer(r"^(\d+)\. (.+)$", MAIN.split("## Working references")[1], re.M)}

PAGE_W, PAGE_H, MARGIN = Cm(21.0), Cm(29.7), Cm(2.5)
TEXT_W = {"portrait": 21.0 - 5.0, "landscape": 29.7 - 5.0}      # cm


# ---- inline LaTeX fragments ---------------------------------------------------------------------------------
def delatex(s):
    """Short inline \\( ... \\) fragments: Greek letters, E_N expressions, relations, signs."""
    def generic(f):
        for a, b in (("\\alpha", "α"), ("\\le", "≤"), ("\\ge", "≥"), ("\\times", "×"), ("\\Delta", "Δ"),
                     ("\\approx", "≈"), ("\\,", " "), ("\\ ", " ")):
            f = f.replace(a, b)
        f = re.sub(r"\\mathrm\{([^}]*)\}", r"\1", f)
        f = f.replace("E_N", "*E*_N").replace("=-", " = −")
        f = re.sub(r"(?<=[\w)])-(?=[\w(])", " − ", f)
        f = re.sub(r"\^\{?([-−]?\w+)\}?", lambda m: "^{" + m.group(1).replace("-", "−") + "}", f)
        return f.replace("=", " = ") if " = " not in f else f
    return re.sub(r"\\\((.*?)\\\)", lambda m: generic(m.group(1)), s)


# ---- citation renumbering (order of first citation in the SI) -----------------------------------------------
CITE = re.compile(r"\[(\d[\d,–\- ]*)\]")


def expand(c):
    out = []
    for part in c.split(","):
        part = part.strip()
        if re.fullmatch(r"\d+[–-]\d+", part):
            a, b = (int(x) for x in re.split("[–-]", part))
            out += list(range(a, b + 1))
        elif part.isdigit():
            out.append(int(part))
    return out


def outside_tables(md):
    """Lines outside markdown tables; citations are renumbered in text and captions, not in table cells."""
    return "\n".join(line for line in md.splitlines() if not line.lstrip().startswith("|"))


order = {}
for c in CITE.findall(outside_tables(NOTES) + "\n" + outside_tables(TABLES)):
    for n in expand(c):
        if n not in order:
            order[n] = len(order) + 1
missing = sorted(set(order) - set(REFS))
assert not missing, f"cited but not in the working references: {missing}"


def compress(nums):
    nums = sorted(set(nums))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(str(nums[i]) if j == i else (f"{nums[i]},{nums[j]}" if j == i + 1 else f"{nums[i]}–{nums[j]}"))
        i = j + 1
    return ",".join(out)


def renumber(s):
    return CITE.sub(lambda m: "^^{" + compress(order[n] for n in expand(m.group(1))) + "}", s)


# ---- run writer: **bold**, *italic*, `code`, [text](url), X_{sub} / X_sub, ^{sup}, ^^{citation}, \* and \_ ---
TOKEN = re.compile(r"(\*\*.+?\*\*|\^\^\{[^}]*\}|\*(?!\s)[^*]+?\*|_\{[^}]*\}|\^\{[^}]*\}|(?<=[A-Za-zεχΣ])_[A-Za-z0-9]+)")
ESC_STAR, ESC_UND = "\ue000", "\ue001"      # escaped \* and \_ while emphasis and subscripts are parsed


def runs(par, text, size=12, bold=False, italic=False):
    text = text.replace("\\*", ESC_STAR).replace("\\_", ESC_UND)
    text = re.sub(r"`([^`]*)`", lambda m: m.group(1).replace("_", ESC_UND).replace("*", ESC_STAR), text)
    text = re.sub(r"\[([^\]]+)\]\((?:https?://|\.{0,2}/)[^)]*\)", r"\1", text)
    for piece in TOKEN.split(text):
        if not piece:
            continue
        b, it, sub, sup = bold, italic, False, False
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            runs(par, piece[2:-2], size, True, italic)
            continue
        if piece.startswith("^^{"):
            piece, sup = piece[3:-1], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            runs(par, piece[1:-1], size, bold, True)
            continue
        elif piece.startswith("_{"):
            piece, sub = piece[2:-1], True
        elif piece.startswith("^{"):
            piece, sup = piece[2:-1], True
        elif piece.startswith("_") and len(piece) > 1:
            m = re.match(r"_([A-Za-z0-9]+)(.*)", piece, re.S)
            if m and m.group(2):
                runs(par, "_{" + m.group(1) + "}", size, bold, italic)
                runs(par, m.group(2), size, bold, italic)
                continue
            piece, sub = piece[1:], True
        r = par.add_run(piece.replace(ESC_STAR, "*").replace(ESC_UND, "_"))
        r.bold, r.italic = b, it
        r.font.size = Pt(size)
        r.font.name = FONT
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        if sub:
            r.font.subscript = True
        if sup:
            r.font.superscript = True


def plain(text):
    """Visible characters of a markdown cell, for width estimates."""
    text = re.sub(r"\\([*_])", r"\1", text)
    return re.sub(r"\*\*|\*|`|\^\^?\{|_\{|\}", "", text)


# ---- document -----------------------------------------------------------------------------------------------
doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = FONT, Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 2.0
st.paragraph_format.space_after = Pt(0)


def page_setup(sec, orient):
    sec.orientation = WD_ORIENT.LANDSCAPE if orient == "landscape" else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = (PAGE_H, PAGE_W) if orient == "landscape" else (PAGE_W, PAGE_H)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, MARGIN)


sec = doc.sections[0]
page_setup(sec, "portrait")
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if tag:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = txt
    r._element.append(el)
orientation = "portrait"


def set_orientation(orient):
    """Start a new section (on a new page) when the orientation changes."""
    global orientation
    if orient != orientation:
        page_setup(doc.add_section(WD_SECTION.NEW_PAGE), orient)
        orientation = orient


def heading(text, level, page_break=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(12 if level <= 2 else 10), Pt(4)
    pf.keep_with_next = True
    pf.page_break_before = page_break
    runs(p, text, size={1: 16, 2: 14, 3: 12}[level], bold=True)
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), str(level - 1))
    p._p.get_or_add_pPr().append(ol)
    return p


def para(text, size=12, indent=True, spacing=2.0, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = spacing
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.6)
    p.paragraph_format.keep_with_next = keep
    runs(p, text, size)
    return p


def list_item(text, marker):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    runs(p, f"{marker}\t" if marker else "", 12)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
    runs(p, text, 12)


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in re.split(r"(?<!\\)\|", line.strip().strip("|"))]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


CHAR_CM = {"portrait": 0.19, "landscape": 0.17}       # mean character width at 9 pt and 8 pt (bold header)


def column_chars(rows):
    """(minimum, wanted) characters per column. Minimum: the longest word, header included, so no word breaks;
    wanted: the longest body cell, damped above 30 characters and capped at 70 so that long text wraps. Both
    include 2 for the cell padding."""
    lo, hi = [], []
    for i in range(len(rows[0])):
        body = [plain(r[i]) for r in rows[1:]] or [""]
        word = max(len(w) for t in [plain(rows[0][i])] + body for w in (t.split() or [""]))
        lo.append(word + 2)
        longest = max(len(t) for t in body)
        hi.append(max(min(longest if longest <= 30 else 30 + 0.4 * (longest - 30), 70), word) + 2)
    return lo, hi


def widths(rows, orient):
    """Column widths (cm) filling the text width: every column gets its minimum, the remaining width is shared in
    proportion to what each column wants beyond it (or to its wanted width when everything fits)."""
    lo, hi = column_chars(rows)
    cw, total = CHAR_CM[orient], TEXT_W[orient]
    if sum(hi) * cw <= total:
        return [total * h / sum(hi) for h in hi]
    base = [x * cw for x in lo]
    extra = [h - x for h, x in zip(hi, lo)]
    room = max(total - sum(base), 0)
    w = [b + (room * e / sum(extra) if sum(extra) else room / len(lo)) for b, e in zip(base, extra)]
    return [total * x / sum(w) for x in w]


def table_orientation(rows):
    """Landscape when the words alone do not fit across a portrait page, or the text would need far more rows."""
    lo, hi = column_chars(rows)
    fits = sum(lo) * CHAR_CM["portrait"] <= TEXT_W["portrait"]
    return "portrait" if fits and sum(hi) * CHAR_CM["portrait"] <= TEXT_W["portrait"] * 1.6 else "landscape"


def set_cell_width(cell, w):
    cell.width = Cm(w)
    tcW = cell._tc.get_or_add_tcPr().get_or_add_tcW()
    tcW.set(qn("w:type"), "dxa")
    tcW.set(qn("w:w"), str(int(w * 567)))


def add_table(rows, orient):
    size = 9 if orient == "portrait" else 8
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    w = widths(rows, orient)
    grid = t._tbl.tblGrid
    for gc, wi in zip(grid.findall(qn("w:gridCol")), w):
        gc.set(qn("w:w"), str(int(wi * 567)))
    numeric = re.compile(r"^[−\-+]?[\d,.]+( ?%)?( \(.*\))?$|^–$|^\d+/\d+$|^yes \(\d+\)$|^no$")
    for i, (row, cells) in enumerate(zip(t.rows, rows)):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        if i == 0:
            hdr = OxmlElement("w:tblHeader")
            trPr.append(hdr)
        for j, (c, text) in enumerate(zip(row.cells, cells)):
            set_cell_width(c, w[j])
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Cm(0)
            filled = [plain(r[j]).strip() for r in rows[1:] if plain(r[j]).strip()]
            col_numeric = bool(filled) and sum(bool(numeric.match(x)) for x in filled) >= 0.75 * len(filled)
            if col_numeric and i > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.keep_with_next = (i == 0)      # the header row stays with the first body row
            runs(p, text, size, bold=(i == 0))
    doc.add_paragraph().paragraph_format.line_spacing = 1.0


# ---- markdown blocks ----------------------------------------------------------------------------------------
def blocks(md):
    """(kind, payload) for headings, paragraphs, list items and tables; comments and rules dropped."""
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    out, para_lines, table_lines = [], [], []

    def flush():
        if para_lines:
            out.append(("p", " ".join(x.strip() for x in para_lines)))
            para_lines.clear()
        if table_lines:
            out.append(("table", parse_table(table_lines)))
            table_lines.clear()

    for line in md.splitlines():
        s = line.strip()
        if s.startswith("|"):
            if para_lines:
                flush()
            table_lines.append(s)
            continue
        if table_lines:
            flush()
        if not s or re.fullmatch(r"-{3,}|\*{3,}", s):
            flush()
        elif m := re.match(r"(#{1,4}) (.+)", s):
            flush()
            out.append((f"h{len(m.group(1))}", m.group(2)))
        elif m := re.match(r"([-*]|\d+\.) (.+)", s):
            flush()
            out.append(("li", (m.group(1), m.group(2))))
        elif para_lines == [] and out and out[-1][0] == "li" and line.startswith("  "):
            kind, (mk, txt) = out[-1]
            out[-1] = (kind, (mk, txt + " " + s))           # continuation of a list item
        else:
            para_lines.append(s)
    flush()
    return out


def render(md, top=False):
    """Write one markdown file. Its leading "# ..." line is dropped for the notes (the document title is written
    once) and becomes a section heading on a new page for the tables (top=True)."""
    items = blocks(md)
    if items and items[0][0] == "h1":
        items = [("h2", items[0][1])] + items[1:] if top else items[1:]
    for k, (kind, x) in enumerate(items):
        if kind in ("h1", "h2", "h3", "h4"):
            level = min(int(kind[1]), 3)
            orient = "portrait"
            if kind == "h3" and x.startswith("Supplementary Table") or kind == "h2" and top:
                # a table heading takes the orientation of its table; the tables' section heading that of the first
                stop = ("h1", "h2") if kind == "h2" else ("h1", "h2", "h3", "h4")
                nxt = next((y for kd, y in items[k + 1:] if kd == "table" or kd in stop), None)
                if isinstance(nxt, list):
                    orient = table_orientation(nxt)
            if orient != orientation:
                set_orientation(orient)
                heading(renumber(delatex(x)), level)
            else:
                heading(renumber(delatex(x)), level, page_break=(kind == "h2" and top))
        elif kind == "p":
            in_table_block = any(kd == "table" for kd, _ in items[k + 1:k + 2])
            if in_table_block or (k and items[k - 1][0] == "h3" and items[k - 1][1].startswith("Supplementary Table")):
                cap = para(renumber(delatex(x)), size=10, indent=False, spacing=1.0, keep=True)
                cap.paragraph_format.space_after = Pt(6)
                cap.paragraph_format.keep_together = True
            else:
                para(renumber(delatex(x)))
        elif kind == "li":
            mk, txt = x
            list_item(renumber(delatex(txt)), "•" if mk in "-*" else mk)
        elif kind == "table":
            add_table(x, orientation)


title = doc.add_paragraph()
runs(title, "**Supplementary Information**", 16)
title.paragraph_format.space_after = Pt(12)
render(NOTES)
set_orientation("portrait")
render(TABLES, top=True)
set_orientation("portrait")

if order:
    heading("Supplementary References", 2, page_break=True)
    for old, new in sorted(order.items(), key=lambda kv: kv[1]):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        runs(p, f"{new}. " + delatex(REFS[old]), 12)

doc.save(args.out)
print(f"wrote {args.out}; {len(order)} references cited: "
      + ", ".join(f"[{o}]→{n}" for o, n in sorted(order.items(), key=lambda kv: kv[1])))
