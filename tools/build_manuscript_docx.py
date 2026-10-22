"""Build a Springer Nature-style Word manuscript from docs/MANUSCRIPT_MAIN_TEXT.md and docs/MAIN_FIGURE_CAPTIONS.md.

Layout: Times New Roman 12 pt, double spacing, continuous line numbers, page numbers, unnumbered headings,
references renumbered in order of first citation (main text, Methods, then figure legends) and set as superscripts,
uncited entries left out, figures with their legends at the end.

    CatalystForge/.venv/python tools/build_manuscript_docx.py OUT.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
TEXT = (ROOT / "docs/MANUSCRIPT_MAIN_TEXT.md").read_text(encoding="utf-8")
CAPS = (ROOT / "docs/MAIN_FIGURE_CAPTIONS.md").read_text(encoding="utf-8")
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Manuscript.docx"
FONT = "Times New Roman"

body_md, refs_md = TEXT.split("## Working references")
REFS = {int(m.group(1)): m.group(2).strip() for m in re.finditer(r"^(\d+)\. (.+)$", refs_md, re.M)}

# ---- LaTeX: inline \( ... \) and display $$ ... $$ become native Word equations (tools/omml.py) -------------
sys.path.insert(0, str(ROOT / "tools"))
from omml import DISPLAY, INLINE, add_display_math, add_math  # noqa: E402


def delatex(s):
    return s.replace("3 × 10^5", "3 × 10^{5}")


# ---- citation renumbering -----------------------------------------------------------------------------------
CITE = re.compile(r"\[(\d[\d,–\- ]*)\]")


def expand(c):
    out = []
    for part in c.split(","):
        part = part.strip()
        if re.fullmatch(r"\d+[–-]\d+", part):
            a, b = (int(x) for x in re.split("[–-]", part))
            out += list(range(a, b + 1))
        elif part.isdigit():
            out.append(int(part))
    return out


order = {}
for c in CITE.findall(body_md + "\n" + CAPS):
    for n in expand(c):
        if n not in order:
            order[n] = len(order) + 1
missing = sorted(set(order) - set(REFS))
assert not missing, missing


def compress(nums):
    nums = sorted(set(nums))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(str(nums[i]) if j == i else (f"{nums[i]},{nums[j]}" if j == i + 1 else f"{nums[i]}–{nums[j]}"))
        i = j + 1
    return ",".join(out)


def renumber(s):
    return CITE.sub(lambda m: "^^{" + compress(order[n] for n in expand(m.group(1))) + "}", s)


# ---- run writer: **bold**, *italic*, X_{sub} / X_sub, ^{sup}, ^^{citation} ---------------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|\^\^\{[^}]*\}|\*(?!\s)[^*]+?\*|_\{[^}]*\}|\^\{[^}]*\}|(?<=[A-Za-zεχΣ])_[A-Za-z0-9]+)")


ESC_STAR = "\ue000"        # stands in for an escaped asterisk (\*) while emphasis is parsed


def runs(par, text, size=12, bold=False, italic=False):
    if "\\(" in text:                                       # inline LaTeX -> native Word equation
        for i, piece in enumerate(INLINE.split(text)):
            if i % 2:
                add_math(par, piece, size)
            elif piece:
                runs(par, piece, size, bold, italic)
        return
    text = text.replace("\\*", ESC_STAR)
    for piece in TOKEN.split(text):
        if not piece:
            continue
        b, it, sub, sup = bold, italic, False, False
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            runs(par, piece[2:-2], size, True, italic)
            continue
        if piece.startswith("^^{"):
            piece, sup = piece[3:-1], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            runs(par, piece[1:-1], size, bold, True)
            continue
        elif piece.startswith("_{"):
            piece, sub = piece[2:-1], True
        elif piece.startswith("^{"):
            piece, sup = piece[2:-1], True
        elif piece.startswith("_") and len(piece) > 1:
            m = re.match(r"_([A-Za-z0-9]+)(.*)", piece, re.S)
            if m and m.group(2):            # subscript right after an italic symbol, e.g. *E*_N
                runs(par, "_{" + m.group(1) + "}", size, bold, italic)
                runs(par, m.group(2), size, bold, italic)
                continue
            piece, sub = piece[1:], True
        r = par.add_run(piece.replace(ESC_STAR, "*"))
        r.bold, r.italic = b, it
        r.font.size = Pt(size)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if sub:
            r.font.subscript = True
        if sup:
            r.font.superscript = True


# ---- document -----------------------------------------------------------------------------------------------
doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = FONT, Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 2.0
st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))
ln = OxmlElement("w:lnNumType")
ln.set(qn("w:countBy"), "1")
ln.set(qn("w:restart"), "continuous")
ln.set(qn("w:distance"), "283")
sec._sectPr.append(ln)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if tag:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = txt
    r._element.append(el)


def heading(text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level <= 2 else 6)
    p.paragraph_format.keep_with_next = True
    runs(p, text, size={1: 16, 2: 13, 3: 12}[level], bold=True)
    pPr = p._p.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), str(level - 1))
    pPr.append(ol)
    return p


def table(lines):
    """A Markdown pipe table (header, separator, rows) as a gridded Word table at 10 pt."""
    cells = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines if not re.match(r"^\|[\s|:-]+\|$", ln.strip())]
    tab = doc.add_table(rows=len(cells), cols=len(cells[0]))
    tab.style = "Table Grid"
    tab.autofit = False
    # column widths (16 cm in total): each column first gets room for its longest word at 10 pt (about 0.2 cm per
    # character plus cell padding); the remainder is shared in proportion to the longest entry of each column
    word = [max(len(w_) for row in cells for w_ in (row[c].replace("*", "").split() or [""])) for c in range(len(cells[0]))]
    entry = [max(len(row[c]) for row in cells) for c in range(len(cells[0]))]
    base = [0.2 * n + 0.45 for n in word]
    spare = max(0.0, 16.0 - sum(base))
    widths = [b + spare * e / sum(entry) for b, e in zip(base, entry)]
    widths = [w * 16.0 / sum(widths) for w in widths]
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tab._tbl.tblPr.append(layout)
    for c, w in enumerate(widths):
        tab.columns[c].width = Cm(w)
        for cell in tab.columns[c].cells:
            cell.width = Cm(w)
    for tr in tab.rows:                                     # a row is never cut across a page break
        tr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for r, row in enumerate(cells):
        for c, val in enumerate(row):
            par = tab.cell(r, c).paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            runs(par, f"**{val}**" if r == 0 else val, 10)
            par.paragraph_format.keep_with_next = r < len(cells) - 1   # the table stays on one page with its caption
    tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph()


def para(text, size=12, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.6)
    runs(p, text, size)
    return p


for block in re.split(r"\n\s*\n", body_md.strip()):
    block = renumber(delatex(block.strip()))
    if not block:
        continue
    if block.startswith("### "):
        heading(block[4:], 3)
    elif block.startswith("## "):
        heading(block[3:], 2)
    elif block.startswith("# "):
        p = heading(block[2:], 1)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para("Author list and affiliations to be confirmed by the corresponding author.", indent=False)
    elif DISPLAY.match(block):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        add_display_math(p, DISPLAY.match(block).group(1), 12)
    elif all(ln.strip().startswith("|") for ln in block.splitlines()):
        table(block.splitlines())
    elif block.startswith("**Table "):
        para(" ".join(block.splitlines()), indent=False).paragraph_format.keep_with_next = True
    else:
        para(" ".join(block.splitlines()), indent=not block[:1].islower())   # lower case: continues after an equation

heading("References", 2)
for old, new in sorted(order.items(), key=lambda kv: kv[1]):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    runs(p, f"{new}. " + delatex(REFS[old]), 12)

# ---- figures and legends ------------------------------------------------------------------------------------
caps = {int(m.group(1)): (m.group(2).strip(), m.group(3).strip())
        for m in re.finditer(r"^## Figure (\d) \| (.+?)\n\n(.+?)(?=\n## |\n---|\Z)", CAPS, re.M | re.S)}
assert sorted(caps) == [1, 2, 3, 4, 5, 6], sorted(caps)
# manuscript figure number -> rendered composite (renumbered 2026-10-07: the field-level figure is Fig. 2)
# reordered 2026-10-15: the methanol mechanism (renderer fig4) follows the field-level figure as Fig. 3
FIGURE_FILES = {1: "fig1/Fig1.png", 2: "fig_field/FigField.png", 3: "fig4/Fig4.png", 4: "fig2/Fig2.png",
                5: "fig_frontier/FigFrontier.png", 6: "fig3/Fig3.png"}   # 2026-10-21: measured-kinetics Fig. 5
for n in range(1, 7):
    pic = doc.add_paragraph()
    pic.paragraph_format.page_break_before = True   # no empty page when the previous one is full
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(ROOT / "figures/composite" / FIGURE_FILES[n]), width=Cm(16.0))
    title, legend = caps[n]
    p = doc.add_paragraph()
    runs(p, f"**Fig. {n} | {renumber(title)}**", 12)
    p = doc.add_paragraph()
    runs(p, renumber(delatex(" ".join(legend.split()))), 12)

# ---- Extended Data figures and legends (figures/extended_data, added 2026-10-07) ----------------------------
ED_CAPS = ROOT / "figures/extended_data/ED_CAPTIONS.md"


def ed_text(s):
    """ED captions use $_{x}$ / $^{x}$ math and `code` paths; map them onto the run grammar."""
    def inner(m):
        t = m.group(1)
        if re.fullmatch(r"[_^][^{].*", t):
            t = t[0] + "{" + t[1:] + "}"
        return t
    s = re.sub(r"\$([^$\n]*?)\$", inner, s)
    s = re.sub(r"\s*Source: .*$", "", s, flags=re.S)   # repository paths stay in ED_CAPTIONS.md and Source Data
    return s.replace("`", "")


if ED_CAPS.exists():
    eds = [(int(m.group(1)), m.group(2).strip(), m.group(3).strip()) for m in
           re.finditer(r"^## Extended Data Fig\. (\d+) \| (.+?)\n\n(.+?)(?=\n## |\Z)", ED_CAPS.read_text(encoding="utf-8"),
                       re.M | re.S)]
    for n, title, legend in eds:
        png = sorted((ROOT / "figures/extended_data").glob(f"EDFig{n}_*.png"))
        pic = doc.add_paragraph()
        pic.paragraph_format.page_break_before = True
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(png[0]), width=Cm(16.0))
        p = doc.add_paragraph()
        runs(p, f"**Extended Data Fig. {n} | {ed_text(title)}**", 12)
        p = doc.add_paragraph()
        runs(p, ed_text(" ".join(legend.split())), 12)

doc.save(OUT)
unused = sorted(set(REFS) - set(order))
print(f"wrote {OUT}; {len(order)} references cited; uncited and left out: {unused}")
