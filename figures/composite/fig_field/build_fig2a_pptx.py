"""Fig. 2a as an editable PowerPoint slide (183 x 74 mm), for final position adjustments by hand.

Text, boxes, arrows, zoom connectors, scale bars and cost bars are native shapes; the rendered structures, the
photograph of extrudates and the converter schematic are pictures. Sources (assets/):
  surface.png, nanoparticle.png   OVITO renders (build_atoms.py)
  extrudates_restored.png         molecular-sieve extrudates, Wikimedia Commons file
                                  "Extrudates-pellets-beads-molecular-sieve-hengye-inc.jpg" (CC BY-SA 4.0), crop
                                  upscaled by an image model without adding or moving objects
  converter.png                   multi-bed ammonia converter, Wikimedia Commons "Ammoniakreaktor MS-ru.svg" (CC0),
                                  labels removed and recoloured
  thumb_*.png                     flat thumbnails from glyphs.py (made here)

    CatalystForge/.venv/python build_fig2a_pptx.py   -> Fig2a.pptx
"""
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from PIL import Image, ImageDraw, ImageFilter  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Mm, Pt  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.dirname(HERE))
import glyphs  # noqa: E402
from style import DARK_B, DARK_G, FE, INK, PALE_B, RED, RU, TINT_B, TINT_G  # noqa: E402

A = os.path.join(HERE, "assets")
W, H = 183.0, 74.0
GOLD = "#E3C16F"
LINE = "#8C929C"
BAND = "#F4F5F7"


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ---------------------------------------------------------------------------------------------- picture assets
def thumbs():
    """Flat thumbnails (transparent PNG, 600 dpi) for the workflow boxes."""
    jobs = {"papers": lambda a, t: glyphs.papers(a, *t, INK, PALE_B, RU, FE),
            "table": lambda a, t: glyphs.table(a, *t, INK, PALE_B, GOLD),
            "selfcheck": lambda a, t: glyphs.selfcheck(a, *t, INK, DARK_G),
            "bound": lambda a, t: glyphs.bound(a, *t, INK, "white", FE),
            "bump": lambda a, t: glyphs.bump(a, *t, INK, PALE_B, RED, DARK_G)}
    plt.rcParams.update({"font.family": "Arial", "mathtext.fontset": "custom", "mathtext.rm": "Arial"})
    for name, fn in jobs.items():
        w, h = 11.0, 13.0
        fig = plt.figure(figsize=(w / 25.4, h / 25.4))
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, w)
        ax.set_ylim(0, h)
        ax.set_aspect("equal")
        ax.axis("off")
        fn(ax, (0.3, 0.3, w - 0.6, h - 0.6))
        fig.savefig(os.path.join(A, "thumb_%s.png" % name), dpi=600, transparent=True)
        plt.close(fig)


def round_photo():
    """The extrudate photograph cut to a disc with a feathered edge (a heap seen through a round window)."""
    im = Image.open(os.path.join(A, "extrudates_restored.png")).convert("RGB")
    s = min(im.size)
    im = im.crop(((im.width - s) // 2, (im.height - s) // 2, (im.width + s) // 2, (im.height + s) // 2))
    m = Image.new("L", (s, s), 0)
    ImageDraw.Draw(m).ellipse((6, 6, s - 6, s - 6), fill=255)
    m = m.filter(ImageFilter.GaussianBlur(3))
    out = im.convert("RGBA")
    out.putalpha(m)
    out.save(os.path.join(A, "extrudates_disc.png"))


def trimmed(name):
    """Crop a transparent PNG to its content; returns path and aspect (w/h)."""
    p = os.path.join(A, name)
    im = Image.open(p).convert("RGBA")
    bb = im.getbbox()
    out = os.path.join(A, "_trim_" + name)
    im.crop(bb).save(out)
    return out, (bb[2] - bb[0]) / (bb[3] - bb[1])


# ---------------------------------------------------------------------------------------------- slide helpers
class Slide:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Mm(W), Mm(H)
        self.s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.sh = self.s.shapes

    def box(self, x, y, w, h, fill=None, line=LINE, lw=0.6, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        b = self.sh.add_shape(shape, Mm(x), Mm(y), Mm(w), Mm(h))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            b.adjustments[0] = radius
        if fill:
            b.fill.solid()
            b.fill.fore_color.rgb = rgb(fill)
        else:
            b.fill.background()
        if line:
            b.line.color.rgb = rgb(line)
            b.line.width = Pt(lw)
        else:
            b.line.fill.background()
        b.shadow.inherit = False
        b.text_frame.text = ""
        return b

    def text(self, x, y, w, h, s, size=6.5, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
             italic=False, fill=None, line=None):
        tb = self.sh.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h)) if fill is None else self.box(x, y, w, h, fill, line,
                                                                                          0.4, 0.3)
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = Mm(0.4)
        tf.margin_top = tf.margin_bottom = Mm(0.1)
        tf.vertical_anchor = anchor
        tf.word_wrap = True
        for i, part in enumerate(s.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            self._runs(p, part, size, bold, color, italic)
        return tb

    @staticmethod
    def _runs(p, s, size, bold, color, italic):
        """'_x' renders x as subscript and '^x' as superscript (single token up to the next space or brace)."""
        import re
        for tok in re.split(r"([_^]\{[^}]*\}|[_^][^\s{]+)", s):
            if not tok:
                continue
            r = p.add_run()
            if tok[0] in "_^":
                r.text = tok[1:].strip("{}")
                r.font._rPr.set("baseline", "-25000" if tok[0] == "_" else "30000")
            else:
                r.text = tok
            r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, "Arial"
            r.font.color.rgb = rgb(color)

    def pic(self, path, x, y, w=None, h=None):
        return self.sh.add_picture(path, Mm(x), Mm(y), Mm(w) if w else None, Mm(h) if h else None)

    def arrow(self, x1, y1, x2, y2, color="#3C4048", lw=0.75, head=True):
        c = self.sh.add_connector(MSO_CONNECTOR.STRAIGHT, Mm(x1), Mm(y1), Mm(x2), Mm(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(lw)
        if head:
            ln = c.line._get_or_add_ln()
            tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "sm", "len": "sm"})
            ln.append(tail)
        return c

    def poly(self, pts, fill, alpha=0.12, line="#7A808A", lw=0.35):
        fb = self.sh.build_freeform(Mm(pts[0][0]), Mm(pts[0][1]), scale=1.0)
        fb.add_line_segments([(Mm(x), Mm(y)) for x, y in pts[1:]], close=True)
        f = fb.convert_to_shape()
        f.fill.solid()
        f.fill.fore_color.rgb = rgb(fill)
        sf = f.fill._xPr.find(qn("a:solidFill"))[0]
        sf.append(sf.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))
        f.line.color.rgb = rgb(line)
        f.line.width = Pt(lw)
        f.shadow.inherit = False
        return f

    def send_back(self, shape):
        tree = self.sh._spTree
        tree.remove(shape._element)
        tree.insert(2, shape._element)


# ---------------------------------------------------------------------------------------------- layout
def build():
    thumbs()
    round_photo()
    S = Slide()
    # top row: literature -> ACSA (three steps) -> leaderboards
    ty, th = 2.5, 18.0
    S.box(2.5, ty, 31, th, "FFFFFF")
    S.pic(os.path.join(A, "thumb_papers.png"), 3.6, ty + 2.2, h=13.6)
    S.text(15.0, ty + 1.2, 18, 4.5, "Literature", 7, True)
    S.text(15.0, ty + 6.0, 18, 10.5, "papers and SI:\n90 methanol\n52 ammonia", 5.6, anchor=MSO_ANCHOR.TOP)
    acsa = S.box(38.5, ty - 1.0, 104.5, th + 2.0, TINT_B, LINE, 0.6, 0.06)
    S.text(40.0, ty - 0.6, 12, 4.2, "ACSA", 7, True, DARK_B)
    steps = [("Extract", "every value with\nunit and page", "language model", RU, "table"),
             ("Self-check", "hand-built results\nreproduced first", "code", DARK_G, "selfcheck"),
             ("Bound", "full optimization\nonly where needed", "code", DARK_G, "bound")]
    for k, (tt, sub, who, wc, g) in enumerate(steps):
        x0 = 44.5 + k * 33.0
        S.box(x0, ty + 3.4, 31.0, th - 4.6, "FFFFFF", None)
        S.pic(os.path.join(A, "thumb_%s.png" % g), x0 + 0.8, ty + 3.8, h=12.6)
        S.text(x0 + 12.4, ty + 4.0, 18, 4.2, tt, 6.8, True)
        S.text(x0 + 12.4, ty + 8.6, 18.4, 7.5, sub, 5.5, anchor=MSO_ANCHOR.TOP)
        S.text(x0 + 31.0 - (15.5 if k == 0 else 8.0), ty + 0.1, 15.5 if k == 0 else 8.0, 2.8, who, 4.8, color="#FFFFFF", align=PP_ALIGN.CENTER, fill=wc)
        if k < 2:
            S.arrow(x0 + 31.2, ty + 10.0, x0 + 32.8, ty + 10.0)
    S.box(148.5, ty, 32, th, "FFFFFF")
    S.pic(os.path.join(A, "thumb_bump.png"), 149.6, ty + 2.2, h=13.6)
    S.text(161.0, ty + 1.2, 19.5, 4.5, "Leaderboards", 7, True)
    S.text(161.0, ty + 6.0, 19.5, 10.5, "laboratory leader\nvs plant-cost\nleader (b, c)", 5.6, anchor=MSO_ANCHOR.TOP)
    S.arrow(33.6, ty + th / 2, 38.4, ty + th / 2)
    # bottom band: the unchanged multiscale chain
    by, bh = 26.5, 46.0
    band = S.box(2.5, by, 178, bh, BAND, LINE, 0.6, 0.04)
    S.arrow(52.0, ty + th + 1.1, 52.0, by - 0.2)
    S.text(53.0, ty + th + 0.6, 22, 4.2, "every candidate", 5.6, italic=True, color=DARK_B)
    S.arrow(164.5, by - 0.2, 164.5, ty + th + 0.1)
    S.text(143.0, ty + th + 0.6, 21, 4.2, "plant cost", 5.6, italic=True, color=DARK_B, align=PP_ALIGN.RIGHT)
    cy = by + 20.0                                          # picture centre line
    pics = {}
    for key, cx, hh in (("surface.png", 18.0, 15.0), ("nanoparticle.png", 47.0, 15.0)):
        p, asp = trimmed(key)
        w_ = hh * asp
        pics[key] = (cx - w_ / 2, cy - hh / 2, w_, hh)
        S.pic(p, *pics[key][:2], w=w_, h=hh)
    d = 19.0
    pics["disc"] = (75.0 - d / 2, cy - d / 2, d, d)
    S.pic(os.path.join(A, "extrudates_disc.png"), *pics["disc"][:2], w=d, h=d)
    p, asp = trimmed("converter.png")
    hh = 28.5
    pics["conv"] = (104.0 - hh * asp / 2, by + 5.6, hh * asp, hh)
    S.pic(p, *pics["conv"][:2], w=hh * asp, h=hh)
    # zoom connectors: each object is a magnified detail of the next, larger one
    for a_, b_, fx, fy in (("surface.png", "nanoparticle.png", 0.45, 0.22), ("nanoparticle.png", "disc", 0.30, 0.45),
                           ("disc", "conv", 0.62, 0.62)):
        x1, y1, w1, h1 = pics[a_]
        x2, y2, w2, h2 = pics[b_]
        sx, sy = x2 + fx * w2, y2 + fy * h2
        r = 1.0
        sq = S.box(sx - r, sy - r, 2 * r, 2 * r, None, "#2B2E33", 0.6, shape=MSO_SHAPE.RECTANGLE)
        f = S.poly([(x1 + w1 - 1.0, y1 + 1.5), (sx - r, sy - r), (sx - r, sy + r), (x1 + w1 - 1.0, y1 + h1 - 1.5)],
                   "#7789B7", 0.10)
        S.send_back(f)
    S.send_back(band)
    S.send_back(acsa)
    # converter -> cost
    S.arrow(116.0, cy, 140.0, cy)
    S.text(116.5, cy - 7.5, 23, 7.0, "flows, duties and\ncatalyst inventory", 5.6, align=PP_ALIGN.CENTER,
           anchor=MSO_ANCHOR.BOTTOM)
    S.text(116.5, cy + 0.8, 23, 7.0, "T, P, separator\nand purge optimized", 5.6, align=PP_ALIGN.CENTER,
           anchor=MSO_ANCHOR.TOP)
    # catalyst-dependent cost, two candidates
    bx, base = 148.0, cy + 8.0
    S.arrow(bx - 2.0, base, bx + 25.0, base, "#2B2E33", 0.6, head=False)
    S.arrow(bx - 2.0, base, bx - 2.0, base - 17.0, "#2B2E33", 0.6, head=False)
    for j, stack in enumerate((((5.0, RU), (3.0, FE), (2.0, GOLD)), ((8.5, RU), (4.5, FE), (2.8, GOLD)))):
        yy = base
        for hh_, c in stack:
            S.box(bx + 3.0 + j * 11.0, yy - hh_, 6.0, hh_, c, "#2B2E33", 0.35, shape=MSO_SHAPE.RECTANGLE)
            yy -= hh_
        S.text(bx + 1.0 + j * 11.0, base + 0.3, 10.0, 3.2, "catalyst %d" % (j + 1), 5.0, align=PP_ALIGN.CENTER)
    # labels and scale bars
    for lab, x, lw_ in (("active site", 18.0, 20), ("supported particle", 47.0, 24), ("catalyst pellets", 75.0, 22),
                        ("multibed converter", 104.0, 22), ("synthesis loop", 128.5, 19), ("cost per tonne", 160.0, 20)):
        S.text(x - lw_ / 2, by + 1.2, lw_, 3.6, lab, 6.0, align=PP_ALIGN.CENTER, fill="FFFFFF", line="#C9CDD3")
    for txt, x in (("~1 nm", 18.0), ("~10 nm", 47.0), ("~5 mm", 75.0), ("~10 m", 104.0)):
        S.arrow(x - 3.0, by + 35.6, x + 3.0, by + 35.0, "#3A3D42", 1.6, head=False)
        S.text(x - 8.0, by + 36.0, 16.0, 3.2, txt, 6.0, True, align=PP_ALIGN.CENTER)
    S.text(148.0, by + 36.0, 26.0, 3.2, "USD or EUR t^{-1}", 6.0, True, align=PP_ALIGN.CENTER)
    for x1, x2, txt, col in ((4.0, 90.0, "kinetics: descriptor turnover or measured rate law", "#D5DAE8"),
                             (92.0, 141.0, "process optimization", "#DCE6D5"),
                             (143.0, 179.0, "catalyst-dependent cost", "#EFE3C2")):
        S.text(x1, by + 40.6, x2 - x1, 4.0, txt, 5.8, align=PP_ALIGN.CENTER, fill=col, line=None)
    out = os.path.join(HERE, "Fig2a.pptx")
    S.prs.save(out)
    print("wrote", out)


if __name__ == "__main__":
    build()
