"""MeOH candidate-state upstream ranking -> economic ranking (rebuilt 2026-09-07 from the frozen D01 v3 workbook).
Reads only MeOH_D01_ExplicitRecycleSeparationEconomics_v3.0.xlsx (Candidate_Inputs, Purge_Sweep @ 2 %, Methanol_Leverage)."""
import csv, hashlib, json, sys
from itertools import combinations
from pathlib import Path
import openpyxl, numpy as np
from scipy.stats import spearmanr, kendalltau
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
WB = Path(sys.argv[1]); OUT = Path(sys.argv[2]); OUT.mkdir(parents=True, exist_ok=True)
wb = openpyxl.load_workbook(WB, data_only=True)
cand = {}
for r in wb["Candidate_Inputs"].iter_rows(min_row=4, values_only=True):
    if r[0] and "Re" in str(r[0]): cand[r[0]] = dict(Re_wt=r[1], T_C=r[2], STY=r[3], X=r[4], S_MeOH=r[5], S_CH4=r[6], S_CO=r[7])
for r in wb["Purge_Sweep"].iter_rows(min_row=4, values_only=True):
    if r[0] in cand and isinstance(r[1], (int, float)) and abs(r[1] - 0.02) < 1e-9:
        cand[r[0]].update(NPC=r[2], recycle_kmol_h=r[3], nonH2CO2_frac=r[4], CH4_frac=r[5], H2_feed=r[6], compression=r[7], EC_MEUR=r[8])
for r in wb["Methanol_Leverage"].iter_rows(min_row=4, values_only=True):
    if r[0] in cand: cand[r[0]].update(L_STY=r[1], L_X=r[2], L_CH4=r[3])
names = list(cand)
for c in cand.values(): c["yield"] = c["X"] * c["S_MeOH"]; c["STY_per_gcat"] = c["STY"] * c["Re_wt"] / 100.0
def ranks(key, reverse=True):
    o = sorted(names, key=lambda n: cand[n][key], reverse=reverse); return {n: i + 1 for i, n in enumerate(o)}
up = {"STY_per_gRe (intrinsic productivity, reaction-case atomic_rank)": ranks("STY"), "single-pass MeOH yield X*S": ranks("yield"), "STY per g catalyst": ranks("STY_per_gcat")}
er = ranks("NPC", reverse=False)
metrics = {}
for lab, ar in up.items():
    a = [ar[n] for n in names]; e = [er[n] for n in names]
    inv = [(x, y) for x, y in combinations(names, 2) if (ar[x] - ar[y]) * (er[x] - er[y]) < 0]
    metrics[lab] = dict(spearman=float(spearmanr(a, e).statistic), kendall=float(kendalltau(a, e).statistic), pairwise_inversions=len(inv), pairs=6,
                        upstream_winner=min(names, key=lambda n: ar[n]), economic_winner=min(names, key=lambda n: er[n]), inverted_pairs=inv)
with open(OUT / "meoh_candidate_ranking_D01v3.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["candidate", "Re_wt%", "T_C", "STY_gMeOH_gRe_h", "X_CO2", "S_MeOH", "S_CH4", "yield_XS", "rank_STY_per_gRe", "rank_yield", "rank_STY_per_gcat", "NPC_EUR_t_2pct_purge", "economic_rank", "recycle_kmol_h", "nonH2CO2_fraction", "CH4_fraction", "H2_feed_EUR_t", "compression_EUR_t", "EC_MEUR", "L_STY", "L_conversion", "L_CH4_suppression"])
    for n in names:
        c = cand[n]; w.writerow([n, c["Re_wt"], c["T_C"], c["STY"], c["X"], c["S_MeOH"], c["S_CH4"], round(c["yield"], 4), up[list(up)[0]][n], up[list(up)[1]][n], up[list(up)[2]][n], round(c["NPC"], 2), er[n], round(c["recycle_kmol_h"], 1), round(c["nonH2CO2_frac"], 4), round(c["CH4_frac"], 4), round(c["H2_feed"], 2), round(c["compression"], 2), round(c["EC_MEUR"], 2), c["L_STY"], c["L_X"], c["L_CH4"]])
prov = {"source_workbook": str(WB), "sha256": hashlib.sha256(WB.read_bytes()).hexdigest(), "sheets": ["Candidate_Inputs", "Purge_Sweep (purge = 0.02)", "Methanol_Leverage"], "catalyst_data": "Gothe et al. ACS Catal. 2025 Table 3 (100 bar, CO2/H2 = 1:4, 500 C prereduction)", "process_anchor": "Processes 2022, 10, 1535; NPC at 2 % purge", "metrics": metrics, "note": "rebuilt 2026-09-07; no dedicated upstream->economic ranking figure existed in the 2026-08-19 v1.0 / 2026-08-22 v2.2 archives (only Figure 11 panel D and the workbook rank-within-T column)"}
(OUT / "meoh_candidate_ranking_D01v3_provenance.json").write_text(json.dumps(prov, indent=1, default=str), encoding="utf-8")
# figure
plt.rcParams.update({"font.size": 13, "axes.labelsize": 14, "axes.titlesize": 14})
short = {n: n.replace(" wt% Re | ", "%Re\n") for n in names}
fig, axes = plt.subplots(1, 3, figsize=(17, 6))
ax = axes[0]; ar = up[list(up)[0]]
for n in names:
    ax.plot(ar[n], er[n], "o", ms=13, mec="black", color="#d62728" if n == metrics[list(up)[0]]["upstream_winner"] else "#1f77b4"); ax.text(ar[n] + 0.12, er[n], short[n].replace("\n", " "), va="center", fontsize=12)
ax.plot([0.5, 4.5], [0.5, 4.5], ":", color="black"); ax.set_xlim(0.5, 5.2); ax.set_ylim(4.6, 0.4); ax.set_xticks([1, 2, 3, 4]); ax.set_yticks([1, 2, 3, 4])
ax.set_xlabel("upstream rank: STY per g Re (intrinsic productivity)"); ax.set_ylabel("economic rank: NPC at 2 % purge")
m = metrics[list(up)[0]]; ax.set_title(f"A  ρ = {m['spearman']:.2f}, τ = {m['kendall']:.2f}, {m['pairwise_inversions']}/6 pairs inverted", fontsize=13); ax.grid(color="#dddddd")
ax = axes[1]; x = np.arange(4); order = sorted(names, key=lambda n: ar[n])
sty = [cand[n]["STY"] for n in order]; npc = [cand[n]["NPC"] for n in order]
ax.bar(x - 0.2, sty, 0.4, color="#2ca02c", edgecolor="black", label="STY (g MeOH / g Re / h)"); ax.set_ylabel("STY (g MeOH / g Re / h)"); ax.set_xticks(x); ax.set_xticklabels([short[n] for n in order])
ax2 = ax.twinx(); ax2.bar(x + 0.2, npc, 0.4, color="#ff7f0e", edgecolor="black", label="NPC (€/t MeOH)"); ax2.set_ylabel("NPC at 2 % purge (€/t MeOH)"); ax2.set_ylim(0, 1400)
for i, v in enumerate(npc): ax2.text(x[i] + 0.2, v + 20, f"{v:.0f}", ha="center", fontsize=11)
ax.set_title("B  candidates in upstream (STY) order", fontsize=13); h1, l1 = ax.get_legend_handles_labels(); h2, l2 = ax2.get_legend_handles_labels(); ax.legend(h1 + h2, l1 + l2, loc="upper left", fontsize=11)
ax = axes[2]
pos = {"1 wt% Re | 200 C": (4, 1035), "1 wt% Re | 250 C": (9, 1000), "5 wt% Re | 200 C": (9, 930), "5 wt% Re | 250 C": (23, 1230)}
for n in names:
    xv, yv = cand[n]["S_CH4"] * 100, cand[n]["NPC"]
    ax.plot(xv, yv, "o", ms=13, mec="black", color="#1f77b4")
    ax.annotate(short[n].replace("\n", " "), (xv, yv), xytext=pos[n], fontsize=12, ha="right" if n.endswith("250 C") and n.startswith("5") else "left", arrowprops=dict(arrowstyle="-", color="black", lw=0.8))
ax.set_xlabel("CH4 selectivity (%)"); ax.set_ylabel("NPC at 2 % purge (€/t MeOH)"); ax.set_title("C  cost follows CH4 selectivity, not STY", fontsize=13); ax.grid(color="#dddddd"); ax.set_xlim(-2, 32)
fig.suptitle("CO2-to-MeOH (D01 v3, source-anchored explicit loop): candidate-state upstream ranking → economic ranking", fontsize=15)
fig.tight_layout(); fig.savefig(OUT / "MeOH_F03_UpstreamToEconomicRanking_D01v3.png", dpi=170); print(json.dumps(metrics, indent=1, default=str))
