"""Recompute every model-written value of the methanol workbook from its Candidate_Inputs sheet.

The workbook keeps the species balance as Excel formulas but stores the economics (Explicit_Loop_2pct columns
Y-AB..AG, Equipment_Breakdown, Purge_Sweep, Engineering_Diagnostic, Methanol_Leverage) as values written by the
original generator. This script recomputes those values with `meoh_d01_model.py`, so the workbook stays a
faithful record of its own inputs.

    python regenerate_d01_values.py --check   # recompute and compare, write nothing
    python regenerate_d01_values.py           # apply INPUT_CORRECTIONS, recompute, write the workbook

Input correction (2026-10-05). Gothe et al., ACS Catal. 15, 19111 (2025), Table 3, row Re 1 wt% / prereduction
500 C / reaction 250 C reports STY 65, CH3OH 97 %, CO 1 %, CH4 1 %, X 23 %. The workbook carried CH4 = 3 %.
Under the workbook's own closure convention (CH4 central value; CO-like = 1 - S_MeOH - S_CH4) the row becomes
S_CH4 = 0.01, S_CO-like = 0.02.
"""
import argparse
import math
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import meoh_d01_model as M  # noqa: E402

WB = HERE / "MeOH_D01_ExplicitRecycleSeparationEconomics_v3.0.xlsx"
INPUT_CORRECTIONS = {  # Candidate_Inputs row -> {column: value}
    7: {"G": 0.01, "H": 0.02, "I": "1%"},
}

ap = argparse.ArgumentParser()
ap.add_argument("--check", action="store_true")
args = ap.parse_args()

wb = openpyxl.load_workbook(WB)  # formulas kept; the stored formula cells carry no cached values
ci = wb["Candidate_Inputs"]
if not args.check:
    for row, cols in INPUT_CORRECTIONS.items():
        for col, v in cols.items():
            ci["%s%d" % (col, row)] = v

cands = []
for r in range(5, 9):
    name, re_wt, t_c, sty, x, smeoh, sch4, sco = (ci.cell(r, j).value for j in range(1, 9))
    c = dict(name=name, Re_wt=float(re_wt), T_C=t_c, STY=float(sty), X=float(x), SMeOH=float(smeoh),
             SCH4=float(sch4), SCO=float(sco))
    assert abs(c["SMeOH"] + c["SCH4"] + c["SCO"] - 1) < 1e-9, name
    cands.append(c)

maxdiff = {}


def put(ws, coord, v):
    old = ws[coord].value
    if isinstance(old, (int, float)) and isinstance(v, (int, float)) and not (isinstance(v, float) and math.isnan(v)):
        maxdiff[ws.title] = max(maxdiff.get(ws.title, 0.0), abs(float(old) - float(v)))
    if not args.check:
        ws[coord] = v


central = {c["name"]: M.cost(c) for c in cands}
sweeps = {c["name"]: M.purge_sweep(c) for c in cands}

# Explicit_Loop_2pct: columns Y..AG (25..33) are stored values
el = wb["Explicit_Loop_2pct"]
for i, c in enumerate(cands):
    e, r = central[c["name"]], 5 + i
    vals = [e["elec_eur_t"], e["catalyst_t"], e["EC_MEUR"], e["sep_loop_EC_MEUR"], e["FCI_MEUR"], e["ACC_MEUR_Y"],
            e["direct_MEUR_Y"], e["indirect_MEUR_Y"], e["NPC_MEUR_Y"]]
    for j, v in enumerate(vals):
        put(el, "%s%d" % (openpyxl.utils.get_column_letter(25 + j), r), float(v))

# Equipment_Breakdown: rows 5..15, candidate columns B..E
eb = wb["Equipment_Breakdown"]
for r in range(5, 16):
    k = eb.cell(r, 1).value
    for j, c in enumerate(cands):
        put(eb, "%s%d" % (openpyxl.utils.get_column_letter(2 + j), r), float(central[c["name"]]["ec"][k]))

# Purge_Sweep: candidate blocks in input order
ps = wb["Purge_Sweep"]
r = 5
for c in cands:
    for d in sweeps[c["name"]]:
        assert ps.cell(r, 1).value == c["name"] and abs(ps.cell(r, 2).value - d["purge"]) < 1e-12, (r, c["name"])
        vals = [d["cost_eur_t"], d["recycle_kmol_h"], d["nonreactive_fraction"], d["methane_fraction"], d["h2_eur_t"],
                d["elec_eur_t"], d["EC_MEUR"]]
        for j, v in enumerate(vals):
            put(ps, "%s%d" % (openpyxl.utils.get_column_letter(3 + j), r), float(v))
        r += 1

# Engineering_Diagnostic
ed = wb["Engineering_Diagnostic"]
for i, c in enumerate(cands):
    e, d, r = central[c["name"]], M.purge_diagnostic(sweeps[c["name"]]), 5 + i
    assert ed.cell(r, 1).value == c["name"]
    for col, v in zip("BCDE", (e["nonreactive_fraction"], e["recycle_kmol_h"], d["purge"], d["cost_eur_t"])):
        put(ed, "%s%d" % (col, r), float(v))
put(ed, "B10", M.SOURCE_NONREACTIVE_REFERENCE)

# Methanol_Leverage
ml = wb["Methanol_Leverage"]
for i, c in enumerate(cands):
    r = 5 + i
    assert ml.cell(r, 1).value == c["name"]
    l_ch4 = M.leverage_ch4_suppression(c)
    for col, v in zip("BCD", (M.leverage_increase(c, "STY"), M.leverage_increase(c, "X"),
                              None if math.isnan(l_ch4) else l_ch4)):
        put(ml, "%s%d" % (col, r), v)

for c in cands:
    print("%-17s S_CH4 %.2f S_CO %.2f  NPC %.6f EUR/t" % (c["name"], c["SCH4"], c["SCO"], central[c["name"]]["cost_eur_t"]))
for k, v in maxdiff.items():
    print("max |new - stored| on %-22s %.3e" % (k, v))
if args.check:
    assert all(v < 1e-6 for v in maxdiff.values()), "workbook values do not reproduce from their inputs"
    print("check passed: every stored value reproduces from Candidate_Inputs")
else:
    wb.save(WB)
    print("wrote", WB.name)
